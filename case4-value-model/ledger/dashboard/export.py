"""X16 part 2: the dashboard's data, exported once from the ledger, the register and the workbook.

Writes `dashboard/data.js` (window.DASH = {...}) for `dashboard/index.html`, a static page that needs no server and
no network. Two layers, never mixed:

- `group`: the programme's figures, read from `assumptions.csv` and `Case4_Value_at_Risk.xlsx`. Every one is logged in
  `trace` with the register row or workbook cell it came from.
- `proto`: the ledger prototype (X4). Real Dutch cars, registration and keeper dates, catalogue prices and the used-car
  price index; synthetic claims, contracts, buyer labels and resale prices (X1). Its totals describe the prototype,
  never the group.

Queues export their top rows, not every car. No VIN, no advert, no person leaves the store. The checks at the end tie
the export back to the store and to the X4 reports; any failure stops the export.

Usage (from case4-value-model/): .venv/bin/python ledger/dashboard/export.py
"""
import csv
import hashlib
import json
import re
import sys
from datetime import date
from pathlib import Path

import duckdb
import numpy as np
import openpyxl
import pandas as pd

HERE = Path(__file__).parent
LEDGER = HERE.parent
MODEL = LEDGER.parent
STORE = MODEL / "data" / "ledger" / "synthetic.duckdb"
PNL = MODEL / "data" / "ledger" / "car_pnl.parquet"
WORKBOOK = MODEL.parent / "Case4_Value_at_Risk.xlsx"
sys.path.insert(0, str(LEDGER))
from readiness import MODEL as READY, WINDOW  # noqa: E402  the queue as it runs now
OUT = HERE / "data.js"
QUEUE_ROWS = 150          # rows per queue on the page; the totals cover every car
REASON_ROWS = 15          # each flag reason's largest claims, so the page can filter its claims queue by reason
SHOWCASE = [51500, 55089, 78]   # X4's showcase cars: the most kinds of event, a duplicate hold, a late-claim hold

# X25's build scenarios as the workbook names them (analysis/programme_cost.py SCENARIOS), and the variant rows of the
# Programme cost sheet that bound each one: what moves its cost, never a fifth scenario
SCENARIOS = ["1. In-house, no AI", "2. Buy and outsource", "3. In-house with AI", "4. Cheapest route per workstream"]
RANGES = {1: [("1 with the coders in the lower-cost hub", "with the coders in a lower-cost hub, no coordination cost counted")],
          2: [("2 at Consip's Italian public team-day", "at an Italian public tender's day rate, a floor"),
              ("2 with the integrator's engineers offshore", "with the integrator's engineers offshore")],
          3: [("3. In-house with AI, upper gain", "with the largest measured AI gain"),
              ("3. In-house with AI, adverse", "with AI slowing experts down")]}

# Each queue's reason types (the tool pass, part 5): a short fixed list, one type per row from a rule on the row's own
# fields, as real queue tools name the type beside every score. The type says what the person does; the row keeps its
# own detail. `{thin}` is the thin-slice switch, filled in at export.
REASON_TYPES = {
    "claims": {
        "Paid twice": "the same car's programme claimed again, in either system: ask the dealer which claim stands",
        "Breaks the terms": "fails a programme rule on its face (filed after the deadline, in the wrong system, or for "
                            "a buyer the programme excludes): the hold stands unless the dealer shows an error",
        "Sale unproven": "the group's own records don't show the sale as claimed (no registered car for the order, no "
                         "fleet contract, or a finance contract that says otherwise): ask for proof of sale",
        "Target gaming": "the dealer's quarter reaches its target only on the claimed dates (registered in its last "
                         "days, by the dealer itself, or on changed orders): ask for the keepers' dates and the order "
                         "history",
    },
    "timing": {
        "48-month lease end": "at the 48-month wave (46 to 49 months), where the Dutch register shows leases ending",
        "60-month lease end": "at the 60-month wave (57 to 60 months), where the Dutch register shows leases ending",
        "Early keeper change": "under two years old: at age one the group's cars change keeper far more than the "
                               "market's",
        "Age curve": "none of these: the chance comes from the car's age alone",
    },
    "desk": {
        "Tied band": "the pricer prices inside the band tied to the engine's uncertainty; beyond it a manager decides",
        "Little history": "fewer than {thin} cars of the model in the book: a wider band, and a person signs off the price",
    },
}
CLAIM_TYPE = {"duplicate": "Paid twice",
              "late_claim": "Breaks the terms", "wrong_system": "Breaks the terms", "ineligible_buyer": "Breaks the terms",
              "unmatched_order": "Sale unproven", "fleet_unverified": "Sale unproven",
              "contract_contradicted": "Sale unproven",
              "gaming_suspect": "Target gaming", "self_registration_suspect": "Target gaming",
              "order_change_suspect": "Target gaming"}

# Each source's as-of date (the tool pass, part 6): what the group's version holds, whether the prototype's is real, and
# the views that read it. The dates come from the store: the latest event and the day the ledger learnt it.
SOURCES = {
    "system_a": ("Claims system A", "incentive claims as filed", "synthetic (X1)", "New-car"),
    "system_b": ("Claims system B", "incentive claims as filed", "synthetic (X1)", "New-car"),
    "x1_layer3": ("Claims controls", "the checks' decision on every claim", "our checks, run on the synthetic claims",
                  "New-car"),
    "vehicle_master": ("Vehicle master", "orders and registrations as recorded", "mixed: real cars and "
                       "first-registration dates (RDW); dealers, buyers and gamed dates synthetic (X1)", "New-car"),
    "oem_prices": ("Price list", "each car's list price", "real catalogue prices (RDW)", "New-car"),
    "oem_record": ("Order record", "the dealer invoiced, after transfers", "synthetic (X1)", "New-car"),
    "oem_transfers": ("Dealer transfers", "cars moved between dealers", "synthetic (X1)", "New-car"),
    "oem_order_log": ("Order log", "orders signed and changed", "synthetic (X1)", "New-car"),
    "register": ("Keeper register, for the claims checks", "the first keeper, the keeper read within a quarter, "
                 "trade-ins", "synthetic (X1)", "New-car"),
    "finance_jv": ("Finance contracts", "the joint ventures' contracts", "synthetic (X1)", "Captive finance"),
    "rdw_register": ("Dutch register (RDW)", "each car's current-keeper date: cars back on the market", "real",
                     "Used-car, Captive finance"),
    "x4_synthetic_resale": ("Resale prices", "the price each returned car sold at", "synthetic (X4)", "Used-car"),
}
ESTIMATES = {   # the two monthly estimate tables, one row per car and month
    "marks": ("Monthly mark", "each car's value and band: catalogue price × the Dutch curve × the second-hand car "
              "index", "model on real inputs", "every view"),
    "readiness": ("Readiness scores", "each car's chance of coming to market in 3 and 12 months",
                  "model on the real register", "Captive finance, Used-car"),
}
QUEUES = ("claims", "upgrade", "incoming", "desk")
# RDW's body type (`inrichting`) on each registration, in English (the tool pass, part 7: body-type icons)
BODY_NAMES = {"hatchback": "hatchback", "MPV": "MPV", "stationwagen": "estate", "sedan": "saloon",
              "kampeerwagen": "camper", "voor rolstoelen toegankelijk voertuig": "wheelchair-accessible",
              "speciale groep": "special purpose", "coupe": "coupé", "lijkwagen": "hearse", "cabriolet": "convertible"}
BODY = {}                 # car id -> body type in English, read once from the store by main()

TRACE = []                # (what, where it came from)
REG = {r["id"]: r for r in csv.DictReader(open(MODEL / "assumptions.csv"))}
WB = openpyxl.load_workbook(WORKBOOK, data_only=True)


def reg(rid, what=None):
    TRACE.append((what or REG[rid]["label"], f"register:{rid} [{REG[rid]['tier']}]"))
    return float(REG[rid]["value"])


def span(vals, d=1):
    """A figure that differs by build scenario, as its range ('11.1–20.1'); one number if they all round the same."""
    lo, hi = f"{min(vals):,.{d}f}", f"{max(vals):,.{d}f}"
    return lo if lo == hi else f"{lo}–{hi}"


def wb(sheet, label, col=2, after=None):
    """Column `col` of the first row whose first cell starts with `label` (below the row starting with `after`)."""
    armed = after is None
    for r in WB[sheet].iter_rows(values_only=True):
        first = r[0].strip() if isinstance(r[0], str) else ""
        if not armed:
            armed = first.startswith(after)
            continue
        if first.startswith(label):
            TRACE.append((f"{sheet}: {label} (col {col})", f"workbook:{sheet}!{label}"))
            return float(r[col - 1])
    raise KeyError(f"{sheet}!{label}")


# ---------------------------------------------------------------------------------------------------- group layer
def group():
    S, B, IP = "Summary", "Benefits", "Internal prices"
    g = {}
    g["var"] = {
        "leaks": [
            {"name": "Leak 1: the spend nobody can see", "short": "Incentives paid wrongly or needlessly",
             "expected": wb(S, "Leak 1 - incentives", 2), "one_in_ten": wb(S, "Leak 1 - incentives", 3)},
            {"name": "Leak 2: the moment nobody catches", "short": "Upgrade moments missed",
             "expected": wb(S, "Leak 2 - upgrade", 2), "one_in_ten": wb(S, "Leak 2 - upgrade", 3)},
            {"name": "Leak 3: the value nobody recovers", "short": "Residual value",
             "expected": wb(S, "Leak 3 - residual value", 2), "one_in_ten": wb(S, "Leak 3 - residual value", 3)},
        ],
        "expected": wb(S, "Total value at risk a year", 2), "one_in_ten": wb(S, "Total value at risk a year", 3),
        "stress": wb(S, "Stress: the whole book", 3),
    }
    g["funnel"] = [
        {"step": "At risk, expected a year", "eur_m": wb(B, "At risk (the headline)", 2)},
        {"step": "Reachable by the levers, every source", "eur_m": wb(B, "Addressable", 2)},
        {"step": "Planned, committed levers", "eur_m": wb(B, "Planned benefit, committed levers", 2)},
        {"step": "The same on large IT projects' record", "eur_m": wb(B, "The same at the reference class", 2)},
        {"step": "Certified by Finance today", "eur_m": wb(B, "Certified today", 2)},
    ]
    g["targeting_upper"] = wb(B, "Upper bound beside it", 2)
    g["levers"] = [{"lever": lab, "eur_m": wb(B, lab, 4), "from": ph} for lab, ph in [
        ("Claims controls", "Phase 2"), ("Retention at the upgrade moment", "Phase 3"),
        ("Resale execution: days cut", "Phase 3"), ("Resale execution: own retail", "Phase 3")]]
    # X25's four build scenarios, side by side: none is chosen (the user, 27 September), so there is no single ask
    PC, SUM = "Programme cost", "Payback by scenario"
    g["programme"] = {"run_cost": reg("run_cost_eur_m_year"),
                      "exit_gate": wb(B, "Gate, every source", 2),
                      "fallback_gate": wb(B, "Gate, on the named fallback", 2), "scenarios": []}
    for i, name in enumerate(SCENARIOS, 1):
        g["programme"]["scenarios"].append({
            "name": name, "build": wb(PC, name, 6, after="The four bases"),
            "by_phase": [wb(PC, name, c, after="By phase") for c in range(2, 6)],
            "range": [[wb(PC, lab, 6, after="The four bases"), what] for lab, what in RANGES.get(i, [])],
            "payback_plan": wb(B, f"{name}, plan", 2, after=SUM), "payback_ref": wb(B, f"{name}, reference class", 2, after=SUM),
            "loss_gate1_ref": wb(B, f"{name}, reference class", 3, after=SUM),
            "value_ref": wb(B, "Value over the short horizon at the group's highest", 2,
                            after=f"Cash by period: {name}, reference class"),
            "floor": wb(B, "Floor: below this", 1 + i), "switch": reg(f"gate_leak_switch_s{i}_pct"),
            "money_gate": reg(f"gate_claims_recovered_s{i}_eur_m")})
    sc = g["programme"]["scenarios"]
    for s in sc:
        s["note"] = ""
    if all(abs(x - y) < 1e-9 for x, y in zip(sc[2]["by_phase"], sc[3]["by_phase"])):
        sc[3]["note"] = "every workstream's cheapest route is in-house with AI, so it costs what 3 does"
    ph_pilot, ph_core = reg("x21_cov_pilot_2024"), reg("x21_cov_core_2024")
    g["phases"] = [
        {"n": 1, "name": "See the spend", "months": "0–6", "where": "Group-wide data; no market live",
         "earns": "A leakage baseline certified by Finance"},
        {"n": 2, "name": "Connect the car", "months": "6–12", "where": f"France, {ph_pilot:.1f}% of EU registrations",
         "earns": "Claims recoveries certified; a measured retention uplift"},
        {"n": 3, "name": "Price the level", "months": "12–18",
         "where": f"France, Italy, Germany, Spain: {ph_core:.1f}%", "earns": "Retention and resale earning; the band decision"},
        {"n": 4, "name": "Close the loop", "months": "18–24", "where": "The core markets, then every market",
         "earns": "Run-rate benefits certified at the exit gate"},
    ]
    g["gates"] = [
        [1, f"Incentive spend linked to a VIN ≥ {reg('gate_vin_link_pct'):.0f}% of claim value; dealer crosswalk ≥ "
            f"{reg('gate_dealer_xwalk_pct'):.0f}%", "Delay"],
        [1, f"Duplicate flags that are real ≥ {reg('gate_dup_precision_pct'):.0f}% (read on "
            f"{reg('x22_audit_n_prec95'):.0f} audited flags)", "Waiver with re-review"],
        [1, f"Real duplicates flagged ≥ {reg('gate_dup_recall_pct'):.0f}% (read on {reg('x22_seed_n_rec80'):.0f} "
            "seeded duplicates)", "Waiver with re-review"],
        [1, f"Leakage baseline certified by Finance (below {span([x['switch'] for x in sc], 2)}%, by build scenario: "
            "start with residual value)", "Back-up"],
        [1, "The monthly mark reconciled against the group's own retailer's realised prices", "Waiver with re-review"],
        [1, "Legal basis agreed for each data source", "Back-up"],
        [2, "The engine recalibrated on the realised prices of backfilled returns", "Delay"],
        [2, f"Ledger within ±{reg('gate_gl_reconcile_pct'):.1f}% of the general ledger every month", "Delay"],
        [2, f"Claims recoveries certified in France ≥ €{span([x['money_gate'] for x in sc])}m a year, by build scenario",
         "Back-up"],
        [2, f"Retention uplift ≥ {reg('gate_uplift_pp'):.0f} points over a randomised control, lower bound above zero",
         "Back-up"],
        [3, f"The engine's 80% band covers 80% ± {reg('gate_engine_coverage_tol_pts'):.0f} points of returned cars "
            f"({reg('x22_engine_cov_n'):.0f} cars)", "Waiver with re-review"],
        [3, f"The engine no worse than the bought guide by more than {reg('gate_engine_margin_pts'):.0f} point "
            f"({reg('x22_engine_pair_n_rho50'):,.0f} paired cars)", "Back-up: the guide keeps the price"],
        [3, "The band decision, from pricers' first proposals logged before the cap", "Not pass or fail"],
        [3, f"Resale execution: ≥ {reg('gate_days_cut_days'):.0f} days cut; routed cars' margin ≥ "
            f"{reg('gate_route_margin_pct'):.0f}% after all costs", "Back-up"],
        [3, f"Certified run rate ≥ the floor, €{span([x['floor'] for x in sc])}m a year by build scenario",
         "Kill: Phase 4 is not funded"],
        [4, f"Run-rate benefits certified ≥ €{g['programme']['exit_gate']:.1f}m a year", "Waiver above the floor; Kill below"],
    ]
    g["kpis_targets"] = {
        "vin_link": reg("gate_vin_link_pct"), "precision": reg("gate_dup_precision_pct"),
        "recall": reg("gate_dup_recall_pct"), "claim_days": reg("kpi_claim_days"),
        "review_rate": reg("kpi_review_rate_pct"), "uplift": reg("gate_uplift_pp"),
        "gl": reg("gate_gl_reconcile_pct"), "band_tol": reg("gate_engine_coverage_tol_pts"),
        "drift": reg("kpi_unseen_drift_pct"), "thin": reg("kpi_thin_error_points"),
        "days_cut": reg("gate_days_cut_days"), "band_cov_heldout": reg("x15_cov_all_asbuilt"),
        "claims_recovered": span([x["money_gate"] for x in sc]), "route_margin": reg("gate_route_margin_pct"),
        "engine_margin": reg("gate_engine_margin_pts")}
    g["prices"] = {
        "p1_low": wb("Prior discounting", "Cost at resale per euro of incentive given, low", 2),
        "p1_high": wb("Prior discounting", "Cost at resale per euro of incentive given, high", 2),
        "p1_book_low": wb("Prior discounting", "Of that, traceable to our own prior discounting - low", 2),
        "p1_book_high": wb("Prior discounting", "Of that, traceable to our own prior discounting - high", 2),
        "p2_option_pct": wb(IP, "Customer-option contract, 48 months", 2),
        "p2_option_eur": reg("x6_price_contract_group_48m"),
        "p2_cold_pts": wb(IP, "plus when the level starts in its cold third", 2),
        "p2_tail_pct": wb(IP, "plus the cost of holding the tail", 2),
        "p2_tail_book_eur_m": wb(IP, "the holding cost on the buy-back book", 2),
        "p2_cold_book_eur_m": wb(IP, "the cold premium on the same book", 2),
        "band_avg_pct": reg("price_band_pct"),
        "day_cost": reg("x17_day_cost_central"), "day_cost_low": reg("x17_day_cost_low"),
        "day_cost_high": reg("x17_day_cost_high"),
        "channel_after": reg("x17_channel_gain_after_costs"), "channel_before": reg("x17_channel_gain_before_opex"),
        "channel_breakeven_days": reg("x17_channel_breakeven_days"),
        "thin_switch_cars": 250}
    TRACE.append(("little-history switch, 250 local cars", "value_engine.py (the pooling result's second step)"))
    g["tied_caps"] = tied_caps()
    g["sample_car"] = {"pltv_low": wb(IP, "Predicted lifetime value at sale", 2),
                       "pltv_high": wb(IP, "Predicted lifetime value at sale", 3),
                       "gain_refinance": wb(IP, "Finance arm, taker refinances with the group", 3),
                       "loss_elsewhere": wb(IP, "Finance arm, taker finances elsewhere", 3),
                       "early_repay_cap": wb(IP, "most it may charge for early repayment", 3),
                       "contribution": reg("margin_per_repeat_sale"),
                       # the waterfall's steps, [measured rate, published rate]
                       "pl": [wb(IP, "Lifetime value per car, the Per car sheet's P&L", c) for c in (2, 3)],
                       "p1_cost": [wb(IP, "less the incentive's resale cost", c) for c in (2, 3)],
                       "p2_hold": [wb(IP, "less the level's price to hold", c) for c in (2, 3)]}
    g["pilot"] = {"per_arm": reg("x15_pilot_blocked_customers"), "gate_pp": reg("gate_uplift_pp"),
                  "interim_extra_pct": reg("x22_obf_extra_pct")}
    g["monthend_benchmark"] = {"other": reg("x2_nl_monthend_other"), "psa": reg("x2_nl_monthend_psa"),
                               "fca": reg("x2_nl_monthend_fca"), "opel": reg("x2_nl_monthend_opel"),
                               "years": "–".join(re.search(r"(\d{4})-(\d{4})$", REG["x2_nl_monthend_other"]["label"]).groups())}
    g["risks"] = [
        ["The finance joint ventures won't share contract data",
         f"The named fallback: cohort timing, the partner contacts its own customers; exit gate €{g['programme']['fallback_gate']:.1f}m"],
        ["Dealers can't or won't evidence sales", "Proof of sale in programme terms; hold, never reject; an appeal"],
        ["The engine learns from asking prices; returns sell at trade",
         "The gap bounded now; recalibrated on realised returns before Phase 3 relies on it"],
        ["A market-level fall during the programme", "Priced and re-marked, never forecast; capital for the tail"],
        ["Benefits over-claimed", "One class per euro; Finance certifies against control groups"],
        ["Escalation of commitment", "Outcomes agreed in advance; outside reviewers; the floor as a named stop"],
    ]
    return g


def tied_caps():
    """The tied override band by age, from the override-band analysis (X13)."""
    text = (MODEL / "analysis" / "override_band_report.md").read_text()
    rows = re.findall(r"^\| (\d+) to (\d+) \| \d+ \| (\d+)% \| ([\d.]+)% \|", text, re.M)
    TRACE.append(("tied override band by age", "analysis/override_band_report.md (X13)"))
    assert len(rows) == 5, rows
    return [{"from": int(a), "to": int(b), "band_half": int(h), "cap": float(c)} for a, b, h, c in rows]


# ---------------------------------------------------------------------------------------------------- action rules
# One action per queue row, from a written rule on the row's own fields: the page never leaves a person to compose it.
def claim_action(checked, sla_days):
    """A flagged claim: a person decides within the service level (working days from the check). Never a reject."""
    due = np.busday_offset(np.datetime64(str(checked)[:10]), int(sla_days), roll="forward")
    return f"A person decides by {due}"   # held or review is the queue's status column


def supply_action(retail, thin):
    """An incoming car: get the sale ready before the car arrives; a thin slice also needs a pricer's sign-off."""
    a = "Pre-price the likely trade-in" if retail else "Plan the lease return"   # the page's rule line says what each means
    return a + ("; a pricer signs off (little history)" if thin else "")


def desk_action(thin):
    """A car back this quarter: a thin slice is signed off by a person; otherwise the band holds and a manager decides
    beyond it. The row's type chip names which; the cap's width stays off the line until the pricer's first proposal
    is typed, so the proposal is the pricer's own (X13)."""
    return "A person signs off the price" if thin else "A manager decides beyond the band"


def timing_type(age):
    """The readiness queues' reason type, on the same ages as the row's reason line."""
    return ("Early keeper change" if age < 24 else "48-month lease end" if 46 <= age <= 49
            else "60-month lease end" if 57 <= age <= 60 else "Age curve")


def desk_type(thin):
    return "Little history" if thin else "Tied band"


def stable(key):
    """A number in [0, 1) from a key's hash: the same on every run and machine, with no seed to keep."""
    return int(hashlib.sha256(key.encode()).hexdigest()[:13], 16) / 16 ** 13


def pilot_arms(cust):
    """X15's design run on the prototype's queue: the group's flag randomised within each dealer, two equal arms (the
    holdout is "not flagged by the group", not "never contacted"). Each dealer's customers, in the order of a stable
    hash of the car number, alternate between the arms, starting with the arm a coin on the dealer's number picks: the
    arms differ by at most one within a dealer and balance overall. `cust` needs `car_id` and `dealer`."""
    h = pd.Series([stable(f"pilot:{c}") for c in cust.car_id], index=cust.index)
    k = h.groupby(cust.dealer).rank(method="first") - 1
    coin = cust.dealer.map(lambda d: stable(f"pilot dealer:{d}") < 0.5)
    return pd.Series(np.where((k % 2 == 0) == coin, "treated", "control"), index=cust.index)


# ---------------------------------------------------------------------------------------------------- prototype layer
def proto(con, g):
    p = {}
    pnl = pd.read_parquet(PNL)
    month = con.execute("select max(month_end) from readiness where model = ?", [READY]).fetchone()[0]
    p["as_of"] = str(month)

    # ---- the book by month, the level and the curve apart
    book = con.execute("""select month_end, count(*) cars, sum(mark_eur::DECIMAL(20, 4))::DOUBLE book,
                                 max("level") lvl, sum((mark_eur / "level")::DECIMAL(20, 4))::DOUBLE at_level_one
                          from marks group by 1 order by 1""").fetchdf()
    p["book"] = [{"m": str(r.month_end)[:7], "cars": int(r.cars), "book": round(r.book), "level": round(r.lvl, 4),
                  "curve": round(r.at_level_one)} for r in book.itertuples()]
    lv = book.lvl.tolist()
    now, past = lv[-1], lv[-37:-1]
    avg36 = sum(past) / 36
    cold_third = now <= sorted(past)[11]          # at or below the 12th of 36 months: the cold third
    p["level_now"] = {"m": str(book.month_end.iloc[-1])[:7], "level": round(now, 4), "avg36": round(avg36, 4),
                      "gap_pct": round(100 * (now / avg36 - 1), 2), "cold_third": bool(cold_third),
                      "option_action": ("Below its three-year average: raise the charge (the re-mark rule)"
                                        if now < avg36 else "At or above its three-year average: hold the charge; "
                                        "never cut it because the level runs hot"),
                      "buyback_action": ("In its cold third: add the cold premium to new buy-back contracts"
                                         if cold_third else "Not in its cold third: no cold premium")}

    # ---- claims (leak 1)
    claims = con.execute("""
        select c.event_id, c.car_id, c.source, c.amount_eur eur, c.valid_date sold, c.recorded_date filed,
               json_extract_string(c.attrs, '$.dealer') dealer, json_extract_string(c.attrs, '$.programme') programme,
               json_extract_string(d.attrs, '$.decision') decision, json_extract_string(d.attrs, '$.reason') reason,
               d.recorded_date checked
        from events c join events d on d.event_type = 'claim_decision'
             and json_extract_string(d.attrs, '$.claim_event') = c.event_id
        where c.event_type = 'claim'""").fetchdf()
    claims["system"] = claims.source.map({"system_a": "A", "system_b": "B"})
    claims["lag"] = (pd.to_datetime(claims.checked) - pd.to_datetime(claims.filed)).dt.days
    flagged = claims[claims.decision != "clear"]
    tool = re.search(r"precision ([\d.]+)%, recall ([\d.]+)%", (LEDGER / "x4_decisions_report.md").read_text())
    p["claims"] = {
        "tool_precision": float(tool.group(1)), "tool_recall": float(tool.group(2)),
        "by_system": [{"system": s, "claims": int(len(d)), "eur": round(d.eur.sum())} for s, d in claims.groupby("system")],
        "linked_share_value": round(100 * claims.loc[claims.car_id.notna(), "eur"].sum() / claims.eur.sum(), 2),
        "decisions": [{"decision": k, "claims": int(len(d)), "eur": round(d.eur.sum())}
                      for k, d in claims.groupby("decision")],
        "reasons": sorted([{"reason": k.replace("_", " "), "decision": d.decision.iloc[0], "claims": int(len(d)),
                            "eur": round(d.eur.sum())} for k, d in flagged.groupby("reason")],
                          key=lambda r: (-r["claims"], r["reason"])),
        "run_date": str(claims.checked.max())[:10],
        "flagged_eur": round(flagged.eur.sum()),
        "lag_median_days": int(flagged.lag.median()),
        "review_share": round(100 * (claims.decision == "review").mean(), 2),
        "total": int(len(claims)), "total_eur": round(claims.eur.sum()),
        # the flow from system to decision, and how long flagged claims waited from filing to check
        "by_system_decision": [{"system": s, "decision": k, "claims": int(len(d)), "eur": round(d.eur.sum())}
                               for (s, k), d in claims.groupby(["system", "decision"])],
        "lag_q": [round(float(x), 1) for x in np.quantile(flagged.lag, np.linspace(0, 1, 101))],
    }
    sla = g["kpis_targets"]["claim_days"]
    day = lambda col: pd.to_datetime(col).values.astype("datetime64[D]")
    p["claims"]["within_sla"] = int((np.busday_count(day(flagged.filed), day(flagged.checked)) <= sla).sum())
    q = flagged.sort_values(["eur", "event_id"], ascending=[False, True]).head(QUEUE_ROWS)

    def qrow(r):
        return {"claim": r.event_id[:8], "car": None if pd.isna(r.car_id) else int(r.car_id),
                "dealer": r.dealer, "programme": r.programme, "system": r.system, "eur": round(r.eur, 2),
                "decision": r.decision, "reason": (r.reason or "").replace("_", " "),
                "reason_type": CLAIM_TYPE.get(r.reason), "filed": str(r.filed)[:10],
                "checked_after_days": int(r.lag), "action": claim_action(r.checked, sla)}
    p["claim_queue"] = [qrow(r) for r in q.itertuples()]
    # each reason's largest claims not already in the queue above (the old page reads only claim_queue)
    more = pd.concat([d.sort_values(["eur", "event_id"], ascending=[False, True]).head(REASON_ROWS)
                      for _, d in flagged.groupby("reason")])
    p["claim_queue_more"] = [qrow(r) for r in more[~more.event_id.isin(q.event_id)].itertuples()]

    # ---- the month-end push, on the prototype's real registration dates
    reg_ = con.execute("""select car_id, valid_date d, json_extract_string(attrs, '$.make') make,
                                 json_extract_string(attrs, '$.buyer') buyer,
                                 json_extract_string(attrs, '$.dealer_id')::INT dealer
                          from events where event_type = 'registration'""").fetchdf()
    d = pd.to_datetime(reg_.d)
    reg_["last3"] = (d.dt.days_in_month - d.dt.day) < 3
    handover = con.execute("select count(*) from events where event_type = 'keeper_handover'").fetchone()[0]
    p["monthend"] = {"share": round(100 * reg_.last3.mean(), 2), "cars": int(len(reg_)),
                     "years": f"{d.dt.year.min()}–{d.dt.year.max()}",
                     "by_make": sorted([{"make": k, "cars": int(len(x)), "share": round(100 * x.last3.mean(), 1)}
                                        for k, x in reg_.groupby("make") if len(x) >= 500],
                                       key=lambda r: (-r["share"], r["make"])),
                     "by_buyer": [{"buyer": k, "cars": int(len(x)), "share": round(100 * x.last3.mean(), 1)}
                                  for k, x in reg_.groupby("buyer")],
                     "handover_3m": int(handover)}

    # ---- the first internal price: the discount's resale cost, on cars the group takes back
    fl = pnl[pnl.contract == "fleet lease"]
    lo, hi = g["prices"]["p1_low"], g["prices"]["p1_high"]
    p["price1"] = {"cars": int(len(fl)), "incentives": round(fl.incentives_paid_eur.sum()),
                   "charge_low": round(fl.incentives_paid_eur.sum() * lo), "charge_high": round(fl.incentives_paid_eur.sum() * hi),
                   "by_make": sorted([{"make": k, "cars": int(len(x)), "incentive_per_car": round(x.incentives_paid_eur.mean()),
                                       "charge_low": round(x.incentives_paid_eur.mean() * lo),
                                       "charge_high": round(x.incentives_paid_eur.mean() * hi)}
                                      for k, x in fl.groupby("make") if len(x) >= 300],
                                     key=lambda r: (-r["cars"], r["make"]))}

    # ---- readiness at the latest month, joined to the mark
    rd = con.execute(f"""select r.car_id, r.age_months, r.p3, r.p3_low, r.p3_high, r.p12, r.retail, r.flag,
                                m.mark_eur, m.low_eur, m.high_eur
                         from readiness r join marks m on m.car_id = r.car_id and m.month_end = r.month_end
                         where r.model = '{READY}' and r.month_end = DATE '{month}'""").fetchdf()
    rd = rd.merge(pnl[["car_id", "make", "model", "contract", "buyer"]], on="car_id", how="left")
    rd["expected_eur"] = rd.p3 * rd.mark_eur
    counts = pnl.groupby(["make", "model"]).size()
    thin = counts[counts < g["prices"]["thin_switch_cars"]].sort_values(kind="stable")
    thin_set = set(thin.index)
    rd["thin"] = [(a, b) in thin_set for a, b in zip(rd.make, rd.model)]

    def reason(age):
        if age < 24:
            return "Early keeper change: at age one the group's cars change keeper far more than the market's"
        if 46 <= age <= 49:
            return "At the 48-month lease-end wave"
        if 57 <= age <= 60:
            return "At the 60-month lease-end wave"
        return f"{age} months old: the age curve"

    def row(r, extra=None):
        out = {"car": int(r.car_id), "make": r.make, "model": r.model, "age": int(r.age_months),
               "p3": round(100 * r.p3, 1), "p3_low": round(100 * r.p3_low, 1), "p3_high": round(100 * r.p3_high, 1),
               "mark": round(r.mark_eur), "low": round(r.low_eur), "high": round(r.high_eur),
               "reason": reason(int(r.age_months)), "reason_type": timing_type(int(r.age_months)), "thin": bool(r.thin),
               "body": BODY[int(r.car_id)]}
        out.update(extra or {})
        return out

    # resale: incoming supply, ranked by expected euros arriving (chance x mark)
    sup = rd.sort_values(["expected_eur", "car_id"], ascending=[False, True])
    p["supply_queue"] = [row(r, {"expected": round(r.expected_eur),
                                 "route": "Fleet: back at lease end" if not r.retail
                                 else "Retail: a trade-in if the customer buys again from the group",
                                 "action": supply_action(bool(r.retail), bool(r.thin))})
                         for r in sup.head(QUEUE_ROWS).itertuples()]
    p["supply_totals"] = [{"book": "retail" if k else "fleet",
                           "cars": int(len(x)), "expected_3m": round(x.p3.sum()), "expected_12m": round(x.p12.sum()),
                           "expected_eur": round(x.expected_eur.sum())} for k, x in rd.groupby("retail")]
    bym = rd.groupby(["make", "model"]).agg(cars=("car_id", "size"), exp3=("p3", "sum"),
                                             eur=("expected_eur", "sum")).reset_index()
    p["supply_by_model"] = [{"make": r.make, "model": r.model, "cars": int(r.cars), "expected_3m": round(r.exp3, 1),
                             "expected_eur": round(r.eur)}
                            for r in bym.sort_values(["eur", "make", "model"], ascending=[False, True, True]).head(20).itertuples()]
    # each of those models' largest expected arrivals not in the queue above, so the page can filter by model
    top_models = {(r["make"], r["model"]) for r in p["supply_by_model"]}
    inq = set(sup.head(QUEUE_ROWS).car_id)
    more = pd.concat([x.head(REASON_ROWS) for k, x in sup.groupby(["make", "model"], sort=False) if k in top_models])
    p["supply_queue_more"] = [row(r, {"expected": round(r.expected_eur),
                                      "route": "Fleet: back at lease end" if not r.retail
                                      else "Retail: a trade-in if the customer buys again from the group",
                                      "action": supply_action(bool(r.retail), bool(r.thin))})
                              for r in more[~more.car_id.isin(inq)].itertuples()]
    p["supply_by_age"] = [{"age": int(a), "cars": int(len(x)), "p3": round(100 * x.p3.iloc[0], 1),
                           "expected_3m": round(x.p3.sum())} for a, x in rd.groupby("age_months") if len(x) >= 100]

    # finance: the upgrade moment, flagged retail customers the captive finances
    up = rd[rd.retail & rd.flag]
    fin = up[up.contract == "private finance"].sort_values(["p3", "mark_eur", "car_id"], ascending=[False, False, True])
    p["upgrade"] = {"window_month": int(reg("window_nl_measured_m", "the upgrade queue's equity window (month)")),
                    "retail_scored": int(rd.retail.sum()), "flagged": int(len(up)),
                    "financed": int(len(fin)), "financed_share": round(100 * len(fin) / len(up), 1)}
    # the pilot (the tool pass, part 5): X15's within-dealer design on the financed customers. Only the treated arm
    # reaches the queues; the control arm leaves the store as a count, so the page cannot show it (hidden by design)
    fin = fin.assign(dealer=fin.car_id.map(reg_.set_index("car_id").dealer))
    allfin = fin.assign(arm=pilot_arms(fin))
    fin = allfin[allfin.arm == "treated"]
    p["upgrade"]["pilot"] = {"treated": int(len(fin)), "control": int((allfin.arm == "control").sum()),
                             "dealers": int(allfin.dealer.nunique())}
    p["upgrade_queue"] = [row(r, {"owner": "The finance partner"}) for r in fin.head(QUEUE_ROWS).itertuples()]
    # the flag is the top tenth by 3-month chance, ties included, so it is a set of ages; each flagged age's largest
    # financed customers not already above, so the page can filter its queue by age (the old page reads only the 150)
    p["upgrade"]["flag_cut"] = round(100 * up.p3.min(), 1)
    p["upgrade"]["flag_ages"] = sorted(int(a) for a in up.age_months.unique())
    # the retention team's calls by contract age: every flagged customer, those the captive finances, and those in the
    # queue (the pilot's treated arm)
    p["upgrade"]["by_age"] = [{"age": a, "flagged": int((up.age_months == a).sum()),
                               "financed": int((allfin.age_months == a).sum()),
                               "queued": int((fin.age_months == a).sum())} for a in p["upgrade"]["flag_ages"]]
    more = pd.concat([x.head(REASON_ROWS) for _, x in fin.groupby("age_months")])
    more = more[~more.car_id.isin(fin.head(QUEUE_ROWS).car_id)]
    p["upgrade_queue_more"] = [row(r, {"owner": "The finance partner"}) for r in more.itertuples()]
    # the readiness curve at every age it has scored (one score per age: an age schedule), the latest month's cars
    curve = con.execute(f"""select age_months a, any_value(p3) p3, any_value(p3_low) lo, any_value(p3_high) hi,
                                     count(distinct round(p3, 9)) n from readiness where model = '{READY}' group by 1
                              order by 1""").fetchdf()
    now_cars = rd.groupby("age_months").size()
    p["ready_curve"] = [{"age": int(r.a), "p3": round(100 * r.p3, 2), "p3_low": round(100 * r.lo, 2),
                         "p3_high": round(100 * r.hi, 2), "cars": int(now_cars.get(r.a, 0)), "scores": int(r.n)}
                        for r in curve.itertuples()]

    # finance: the level result on leases that came back (real index and dates, synthetic residuals)
    back = pnl[pnl.came_back.notna() & pnl.level_part_eur.notna()].copy()
    back["year"] = pd.to_datetime(back.came_back).dt.year
    p["level_result"] = [{"year": int(y), "cars": int(len(x)), "level_part": round(x.level_part_eur.sum()),
                          "per_car": round(x.level_part_eur.mean())} for y, x in back.groupby("year")]
    p["level_result_total"] = {"cars": int(len(back)), "level_part": round(back.level_part_eur.sum()),
                               "residual_set": round(back.residual_set_eur.sum())}

    # resale: cars that came back, by month (real dates), and a pricing desk for the latest quarter
    ctm = con.execute("""select date_trunc('month', valid_date) m, count(*) n from events
                         where event_type = 'came_to_market' group by 1 order by 1""").fetchdf()
    p["returns_by_month"] = [{"m": str(r.m)[:7], "cars": int(r.n)} for r in ctm.itertuples()]
    caps = g["tied_caps"]

    def cap(age_y):
        for c in caps:
            if c["from"] <= age_y < c["to"]:
                return c["cap"]
        return caps[-1]["cap"]

    desk = pnl[pnl.came_back.notna()].copy()
    desk["cb"] = pd.to_datetime(desk.came_back)
    desk = desk[(desk.cb >= pd.Timestamp(month) - pd.offsets.MonthBegin(3)) & (desk.cb <= pd.Timestamp(month))]
    desk["age_y"] = (desk.cb - pd.to_datetime(desk.registration_date)).dt.days / 365.25
    # the desk's day (the tool pass): every car back this quarter, and those past in-house retail's break-even days
    days_back = (pd.Timestamp(month) - desk.cb).dt.days
    p["desk_count"] = int(len(desk))
    p["desk_past_breakeven"] = int((days_back > g["prices"]["channel_breakeven_days"]).sum())
    desk = desk.sort_values(["mark_now_eur", "car_id"], ascending=[False, True]).head(60)
    dc = g["prices"]["day_cost"]
    p["pricing_desk"] = [{"car": int(r.car_id), "make": r.make, "model": r.model, "came_back": str(r.cb.date()),
                          "age": round(r.age_y, 1), "mark": round(r.mark_now_eur), "low": round(r.mark_now_low_eur),
                          "high": round(r.mark_now_high_eur), "cap_pct": cap(r.age_y),
                          "day_cost": round(r.mark_now_eur * dc / 10000, 2), "thin": (r.make, r.model) in thin_set,
                          "reason_type": desk_type((r.make, r.model) in thin_set), "body": BODY[int(r.car_id)],
                          "action": desk_action((r.make, r.model) in thin_set)} for r in desk.itertuples()]
    p["reason_types"] = {q: [[t, m.format(thin=g["prices"]["thin_switch_cars"])] for t, m in d.items()]
                         for q, d in REASON_TYPES.items()}
    p["thin"] = {"models": int(len(thin)), "cars": int(thin.sum()), "of_models": int(len(counts)),
                 "list": [{"make": k[0], "model": k[1], "cars": int(v)} for k, v in thin.head(25).items()],
                 "counts": sorted(int(v) for v in counts.values)}

    # ---- the operating metric: what the ledger fills today
    p["metric_lines"] = [
        ["Incentives paid, held, in review", "Filled (synthetic claims)"],
        ["The discount's resale cost (first internal price)", "Computed on fleet leases: incentive × the measured range"],
        ["The level's price (second internal price)", "Not yet per car: needs each contract's residual and term "
                                                      "(the joint ventures' data); the rule runs on the index today"],
        ["The car's value and band (the engine)", "Filled monthly (real catalogue prices and index)"],
        ["New-car margin", "Not held: the group's data"], ["Finance income", "Not held: the joint ventures' data"],
        ["Remarketing cost", "Not held: the group's data"], ["Contract term and balance", "Not held: the joint ventures' data"],
        ["Battery health, for electric cars", "Not held: readable by the car's legal purchaser from 18 August 2024 "
                                              "(Batteries Regulation, Art. 14); the prototype holds no fuel type"],
    ]
    reg("battreg_soh_access_from", "battery health readable by the legal purchaser (metric line)")
    p["book_now"] = p["book"][-1]
    p["cars_total"] = int(len(pnl))
    return p, claims, rd, allfin[["car_id", "dealer", "arm"]]


# ---------------------------------------------------------------------------------------------------- search, sources
def bodies(con):
    """Each car's body type, from its registration event, in English."""
    rows = con.execute("""select car_id, json_extract_string(attrs, '$.body') from events
                          where event_type = 'registration'""").fetchall()
    return {int(c): BODY_NAMES.get(b, b) for c, b in rows}


def listed(p):
    """The cars each queue lists (both of its parts), by queue."""
    rows = {"claims": p["claim_queue"] + p["claim_queue_more"], "upgrade": p["upgrade_queue"] + p["upgrade_queue_more"],
            "incoming": p["supply_queue"] + p["supply_queue_more"], "desk": p["pricing_desk"]}
    return {q: {r["car"] for r in rows[q] if r["car"] is not None} for q in QUEUES}


def car_index(p):
    """Every car the page holds a record for (the queues' cars and X4's showcase cars), for the search by number or
    model, with the queues that list it. The export carries no VIN, by design, so the car number is the key."""
    info, where = pd.read_parquet(PNL).set_index("car_id"), listed(p)
    return [{"car": int(c), "make": info.make[c], "model": info.model[c], "body": BODY[int(c)],
             "in": [q for q in QUEUES if c in where[q]]}
            for c in sorted(set().union(*where.values(), SHOWCASE))]


def source_dates(con):
    """Each source's as-of date: its latest event and the day the ledger learnt it (every event carries both)."""
    ev = con.execute("""select source, count(*) n, max(valid_date) hap, max(recorded_date) lrn
                        from events group by 1""").fetchdf().set_index("source")
    out = [{"source": s, "name": n, "holds": h, "kind": k, "views": v, "records": int(ev.n[s]),
            "happened": str(ev.hap[s])[:10], "learnt": str(ev.lrn[s])[:10]} for s, (n, h, k, v) in SOURCES.items()]
    for t, (n, h, k, v) in ESTIMATES.items():
        only = f"where model = '{READY}'" if t == "readiness" else ""     # the engine the queues run on
        c, hap, lrn = con.execute(f"select count(*), max(month_end), max(recorded_date) from {t} {only}").fetchone()
        out.append({"source": t, "name": n, "holds": h, "kind": k, "views": v, "records": int(c),
                    "happened": str(hap)[:10], "learnt": str(lrn)[:10]})
    return out


# ---------------------------------------------------------------------------------------------------- one car
def cars(con, ids):
    ids = sorted(set(int(i) for i in ids if i is not None))
    idl = ",".join(map(str, ids))
    ev = con.execute(f"""select car_id, valid_date, recorded_date, event_type, source, amount_eur, kind, attrs
                         from events where car_id in ({idl}) order by car_id, valid_date, recorded_date, seq""").fetchdf()
    dec = con.execute(f"""select json_extract_string(d.attrs, '$.claim_event') ce, d.recorded_date,
                                 json_extract_string(d.attrs, '$.decision') decision,
                                 json_extract_string(d.attrs, '$.reason') reason, c.car_id, c.valid_date
                          from events d join events c on c.event_id = json_extract_string(d.attrs, '$.claim_event')
                          where d.event_type = 'claim_decision' and c.car_id in ({idl})
                          order by c.car_id, c.valid_date, d.recorded_date, ce""").fetchdf()
    mk = con.execute(f"""select car_id, month_end, mark_eur, low_eur, high_eur from marks where car_id in ({idl})
                         order by 1, 2""").fetchdf()
    rs = con.execute(f"""select car_id, month_end, p3, p12 from readiness where model = '{READY}' and car_id in ({idl})
                         order by 1, 2""").fetchdf()
    pnl = pd.read_parquet(PNL).set_index("car_id")
    keep = {"dealer", "programme", "contract_type", "event", "from_dealer", "to_dealer", "first_keeper", "buyer",
            "make", "model", "body", "export", "days_after_registration"}
    out = {}
    for cid in ids:
        e = ev[ev.car_id == cid]
        rows = []
        for r in e.itertuples():
            a = json.loads(r.attrs) if isinstance(r.attrs, str) else {}
            detail = ", ".join(f"{k.replace('_', ' ')}: {v}" for k, v in a.items() if k in keep and v not in (None, ""))
            rows.append([str(r.valid_date)[:10], str(r.recorded_date)[:10], r.event_type.replace("_", " "), r.source,
                         None if pd.isna(r.amount_eur) else round(r.amount_eur, 2), r.kind, detail])
        for r in dec[dec.car_id == cid].itertuples():
            if r.decision != "clear":
                rows.append([str(r.valid_date)[:10], str(r.recorded_date)[:10], f"claim {r.decision}", "claims controls", None,
                             "decision", (r.reason or "").replace("_", " ")])
        rows.sort(key=lambda x: (x[0], x[1]))
        info = pnl.loc[cid]
        out[str(cid)] = {
            "make": info.make, "model": info.model, "body": BODY[cid], "registered": str(info.registration_date.date()),
            "catalogue": round(info.catalogue_eur), "buyer": info.buyer,
            "contract": None if pd.isna(info.contract) else info.contract,
            "incentives": round(info.incentives_paid_eur or 0),
            "events": rows,
            "marks": [[str(r.month_end)[:7], round(r.mark_eur), round(r.low_eur), round(r.high_eur)]
                      for r in mk[mk.car_id == cid].itertuples()],
            "ready": [[str(r.month_end)[:7], round(100 * r.p3, 1), round(100 * r.p12, 1)]
                      for r in rs[rs.car_id == cid].itertuples()],
        }
    return out


# ---------------------------------------------------------------------------------------------------- checks
def checks(con, g, p, claims, rd, arms, blob):
    res = []

    def ok(name, cond, detail=""):
        res.append((name, bool(cond), detail))

    pnl = pd.read_parquet(PNL)
    ok("the book's latest total equals the per-car marks in car_pnl",
       abs(p["book_now"]["book"] - pnl.mark_now_eur.sum()) < 1, f"{p['book_now']['book']:,} vs {pnl.mark_now_eur.sum():,.0f}")
    b = p["book_now"]
    moved = con.execute('select sum(mark_eur / "level" * ("level" * 0.9)) from marks '
                        'where month_end = (select max(month_end) from marks)').fetchone()[0]
    ok("a −10% level move takes the book to 90% of itself (the page's slider does the same)",
       abs(moved - 0.9 * b["book"]) < 1, f"{moved:,.0f}")
    ok("book = curve value × level, every month",
       all(abs(r["book"] - r["curve"] * r["level"]) <= 0.01 * r["book"] + 1 for r in p["book"]))
    store = (LEDGER / "x4_store_report.md").read_text()
    for src, s in (("system_a", "A"), ("system_b", "B")):
        m = re.search(rf"\| {src} \| claim \| ([\d,]+) \| ([\d,]+) \|", store)
        exp_n, exp_eur = int(m.group(1).replace(",", "")), int(m.group(2).replace(",", ""))
        got = next(x for x in p["claims"]["by_system"] if x["system"] == s)
        ok(f"system {s} claims equal the store report", got["claims"] == exp_n and abs(got["eur"] - exp_eur) <= 1,
           f"{got['claims']:,} / €{got['eur']:,}")
    dec = (LEDGER / "x4_decisions_report.md").read_text()
    for k in ("clear", "hold", "review"):
        m = re.search(rf"\| {k} \| ([\d,]+) \|", dec)
        got = next(x for x in p["claims"]["decisions"] if x["decision"] == k)
        ok(f"{k} decisions equal the decisions report", got["claims"] == int(m.group(1).replace(",", "")))
    c = p["claims"]
    ok("the flow from system to decision adds up to both the system and the decision totals",
       all(sum(x["claims"] for x in c["by_system_decision"] if x[key] == t[key]) == t["claims"]
           and abs(sum(x["eur"] for x in c["by_system_decision"] if x[key] == t[key]) - t["eur"]) <= 3   # rounded per flow
           for key, tot in (("system", c["by_system"]), ("decision", c["decisions"])) for t in tot))
    rows = p["claim_queue"] + p["claim_queue_more"]
    ok("every flag reason has its largest claims in the queues (up to 15 each), none twice",
       all(sum(r["reason"] == x["reason"] for r in rows) >= min(REASON_ROWS, x["claims"]) for x in c["reasons"])
       and len({r["claim"] for r in rows}) == len(rows))
    ok("the wait's quantiles hold the reported median; the count within the service level is counted in working days",
       int(c["lag_q"][50]) == c["lag_median_days"] and c["lag_q"] == sorted(c["lag_q"])
       and 0 <= c["within_sla"] <= sum(x["claims"] for x in c["decisions"] if x["decision"] != "clear"),
       f"median {c['lag_q'][50]:,.0f} days; {c['within_sla']:,} within {g['kpis_targets']['claim_days']:.0f} working days")
    u, rc = p["upgrade"], p["ready_curve"]
    ok("the readiness curve has one score per age, and matches the supply by age where both exist",
       all(x["scores"] == 1 for x in rc)
       and all(abs(next(x["p3"] for x in rc if x["age"] == a["age"]) - a["p3"]) < 0.06 for a in p["supply_by_age"]))
    ok("the flagged ages are exactly the scored ages at or above the flag's cut",
       u["flag_ages"] == [x["age"] for x in rc if x["cars"] and round(x["p3"], 1) >= u["flag_cut"]],
       f"ages {u['flag_ages']} at or above {u['flag_cut']}%")
    ok("the calls by contract age add up to the flagged and the financed, age by age",
       [x["age"] for x in u["by_age"]] == u["flag_ages"] and sum(x["flagged"] for x in u["by_age"]) == u["flagged"]
       and sum(x["financed"] for x in u["by_age"]) == u["financed"]
       and all(x["financed"] <= x["flagged"] for x in u["by_age"]),
       ", ".join(f"{x['age']} months {x['financed']:,} of {x['flagged']:,}" for x in u["by_age"]))
    be = g["prices"]["channel_breakeven_days"]
    listed_past = sum((pd.Timestamp(p["as_of"]) - pd.Timestamp(r["came_back"])).days > be for r in p["pricing_desk"])
    ok("the desk's counts hold its listed cars: back this quarter, and past the break-even days",
       p["desk_count"] >= len(p["pricing_desk"]) and listed_past <= p["desk_past_breakeven"] <= p["desk_count"],
       f"{p['desk_count']:,} back, {p['desk_past_breakeven']:,} past {be:.0f} days ({listed_past} of the "
       f"{len(p['pricing_desk'])} listed)")
    uq = p["upgrade_queue"] + p["upgrade_queue_more"]
    ok("every flagged age has its treated customers in the queues (up to 15 each), none twice",
       all(sum(r["age"] == a for r in uq) >= 1 for a in u["flag_ages"]) and len({r["car"] for r in uq}) == len(uq)
       and all(r["age"] in u["flag_ages"] for r in uq))
    pl = u["pilot"]
    lean = arms.groupby("dealer").arm.agg(lambda a: int((a == "treated").sum() - (a == "control").sum()))
    again = pilot_arms(arms.sample(frac=1, random_state=0)).reindex(arms.index)
    ok("the pilot's arms: every financed customer in one arm, within one of each other at every dealer, the same "
       "whatever order the customers come in",
       arms.car_id.is_unique and len(arms) == u["financed"] and arms.dealer.notna().all() and lean.abs().max() <= 1
       and (again == arms.arm).all() and pl["treated"] + pl["control"] == u["financed"],
       f"{pl['treated']:,} treated, {pl['control']:,} held out, {pl['dealers']} dealers, "
       f"{int((lean != 0).sum())} with one more in an arm")
    held = set(arms.car_id[arms.arm == "control"])
    ok("no held-out customer in the upgrade queues; the queued counts by age add up to the treated arm",
       not held & {r["car"] for r in uq} and sum(x["queued"] for x in u["by_age"]) == pl["treated"]
       and all(x["queued"] <= x["financed"] for x in u["by_age"]),
       ", ".join(f"{x['age']} months {x['queued']:,} of {x['financed']:,}" for x in u["by_age"]))
    sq = p["supply_queue"] + p["supply_queue_more"]
    T = {q: [t for t, _ in d] for q, d in p["reason_types"].items()}
    flagged_reasons = set(claims.reason[claims.decision != "clear"])
    used = {q: sorted({r["reason_type"] for r in rows}) for q, rows in
            (("claims", p["claim_queue"] + p["claim_queue_more"]), ("upgrade", uq), ("incoming", sq),
             ("desk", p["pricing_desk"]))}
    ok("every queue row carries one reason type from its queue's fixed list, the one its rule gives; every flagged "
       "claim's reason has a type",
       flagged_reasons <= set(CLAIM_TYPE) and set(CLAIM_TYPE.values()) == set(T["claims"])
       and all(r["reason_type"] == CLAIM_TYPE[r["reason"].replace(" ", "_")] for r in p["claim_queue"] + p["claim_queue_more"])
       and all(r["reason_type"] == timing_type(r["age"]) and r["reason_type"] in T["timing"] for r in uq + sq)
       and all(r["reason_type"] == desk_type(r["thin"]) and r["reason_type"] in T["desk"] for r in p["pricing_desk"]),
       "; ".join(f"{q}: {', '.join(v)}" for q, v in used.items()))
    ok("each of the top models by expected euros has its largest arrivals in the queues (up to 15 each), none twice",
       all(sum(r["make"] == m["make"] and r["model"] == m["model"] for r in sq) >= min(REASON_ROWS, m["cars"])
           for m in p["supply_by_model"]) and len({r["car"] for r in sq}) == len(sq))
    th = p["thin"]
    ok("every model's count is exported; those under the switch are the little-history models",
       len(th["counts"]) == th["of_models"] and sum(1 for c in th["counts"] if c < g["prices"]["thin_switch_cars"]) == th["models"]
       and sum(c for c in th["counts"] if c < g["prices"]["thin_switch_cars"]) == th["cars"])
    rr = (LEDGER / "x4_readiness_report.md").read_text()
    m = re.search(r"Flagged \(top tenth by 3-month chance, ties included\): ([\d,]+)", rr)
    ok("flagged customers equal the readiness report", p["upgrade"]["flagged"] == int(m.group(1).replace(",", "")))
    m = re.search(r"captive finances \(SYNTHETIC share.*?\): ([\d.]+)%", rr)
    ok("the financed share equals the readiness report", abs(p["upgrade"]["financed_share"] - float(m.group(1))) < 0.05)
    for book, n3 in re.findall(r"^\| (retail|fleet) \| [\d,]+ \| ([\d,]+) \|", rr, re.M):
        got = next(x for x in p["supply_totals"] if x["book"] == book)
        ok(f"{book}: cars expected within 3 months equal the readiness report",
           got["expected_3m"] == int(n3.replace(",", "")), f"{got['expected_3m']:,}")
    pr = (LEDGER / "x4_pnl_report.md").read_text()
    m = re.search(r"\*\*Level part €([\d.]+)m\*\*", pr)
    ok("the level result equals the P&L report", abs(p["level_result_total"]["level_part"] / 1e6 - float(m.group(1))) < 0.05)
    hdr = [r for r in WB["Benefits"].iter_rows(values_only=True) if r[0] == "gate"]
    ok("the Benefits sheet's gate columns are the scenarios, in order (the floor is read by column)",
       len(hdr) == 1 and list(hdr[0][1:5]) == SCENARIOS)
    for s in g["programme"]["scenarios"]:
        n = s["name"]
        ok(f"{n}: payback and first-gate loss match the scenario's own case blocks",
           abs(s["payback_plan"] - wb("Benefits", "Payback month", 2, after=f"Cash by period: {n}, plan")) < 1e-6
           and abs(s["payback_ref"] - wb("Benefits", "Payback month", 2, after=f"Cash by period: {n}, reference class")) < 1e-6
           and abs(s["loss_gate1_ref"] - wb("Benefits", "Loss if stopped at the Phase 1 gate", 2,
                                            after=f"Cash by period: {n}, reference class")) < 1e-6)
        ok(f"{n}: the phase costs add up to the build, and the Benefits sheet's plan pays them",
           abs(sum(s["by_phase"]) - s["build"]) < 1e-6
           and all(abs(x - wb("Benefits", "build cost", c, after=f"Cash by period: {n}, plan")) < 1e-6
                   for c, x in enumerate(s["by_phase"], 2)), f"€{s['build']:.1f}m")
    sla = g["kpis_targets"]["claim_days"]
    chk = dict(zip(claims.event_id.str[:8], claims.checked))
    ok("every claim, incoming-stock and pricing-desk row carries the one action its rule gives",
       all(r["action"] == claim_action(chk[r["claim"]], sla) for r in p["claim_queue"] + p["claim_queue_more"])
       and all(r["action"] == supply_action(r["route"].startswith("Retail"), r["thin"]) for r in p["supply_queue"])
       and all(r["action"] == desk_action(r["thin"]) for r in p["pricing_desk"]))
    counts = pnl.groupby(["make", "model"]).size()
    ok("little-history flags follow the switch (fewer than 250 cars of the model in the book)",
       all(r["thin"] == (counts[(r["make"], r["model"])] < g["prices"]["thin_switch_cars"])
           for q in ("supply_queue", "upgrade_queue", "pricing_desk") for r in p[q]))
    ok("the finance queue holds no car under the equity window", all(r["age"] >= WINDOW for r in p["upgrade_queue"]))
    sc = g["sample_car"]
    ok("the sample car's waterfall adds up: its P&L less the two charges is its predicted lifetime value, at both rates",
       all(abs(sc["pl"][i] + sc["p1_cost"][i] + sc["p2_hold"][i] - sc[k]) < 0.01
           for i, k in enumerate(("pltv_low", "pltv_high"))), f"€{sc['pltv_low']:,.2f} and €{sc['pltv_high']:,.2f}")
    ci, where, info = p["car_index"], listed(p), pnl.set_index("car_id")
    ok("search: every listed car and showcase car is in the index once, with its make and model and the queues that "
       "list it; the index is exactly the cars whose record the page holds",
       [x["car"] for x in ci] == sorted(set().union(*where.values(), SHOWCASE))
       and sorted(int(k) for k in p["cars"]) == [x["car"] for x in ci]
       and all(x["make"] == info.make[x["car"]] and x["model"] == info.model[x["car"]] for x in ci)
       and all(x["in"] == [q for q in QUEUES if x["car"] in where[q]] for x in ci)
       and all(str(x["car"]) in p["cars"] for x in ci),
       f"{len(ci):,} cars indexed; {len(p['cars']):,} records in the drawer")
    raw = {b for (b,) in con.execute("""select distinct json_extract_string(attrs, '$.body') from events
                                      where event_type = 'registration'""").fetchall()}
    bodied = p["upgrade_queue"] + p["upgrade_queue_more"] + p["supply_queue"] + p["supply_queue_more"] + \
        p["pricing_desk"] + ci
    ok("body types: every body type in the store has an English name; every car in the queues, the search and the "
       "drawer carries its registration's",
       raw <= set(BODY_NAMES) and all(r["body"] == BODY[r["car"]] for r in bodied)
       and all(c["body"] == BODY[int(k)] for k, c in p["cars"].items()),
       ", ".join(f"{n} {sum(1 for b in BODY.values() if b == n):,}" for n in dict.fromkeys(BODY_NAMES.values())))
    S = {x["source"]: x for x in p["sources"]}
    stored = {r[0] for r in con.execute("select distinct source from events").fetchall()}
    early = sum(con.execute(q).fetchone()[0] for q in (
        "select count(*) from events where recorded_date < valid_date",
        "select count(*) from marks where recorded_date < month_end",
        "select count(*) from readiness where recorded_date < month_end"))
    reported = {}
    for s, n in re.findall(r"^\| (\w+) \| \w+ \| ([\d,]+) \|", store, re.M):
        reported[s] = reported.get(s, 0) + int(n.replace(",", ""))
    num = lambda pat, text: int(re.search(pat, text).group(1).replace(",", ""))
    ret = (LEDGER / "x4_returns_report.md").read_text()
    back = num(r"Came back to market \(latest change after 3 months\): ([\d,]+)", ret)
    reported["rdw_register"] = num(r"Handover within 3 months: ([\d,]+)", ret) + back
    reported["x4_synthetic_resale"] = back
    reported["x1_layer3"] = sum(num(rf"\| {k} \| ([\d,]+) \|", dec) for k in ("clear", "hold", "review"))
    read_on = re.search(r"Register read on (\d{4}-\d{2}-\d{2})", ret).group(1)
    ok("each source's as-of date: every source in the store listed once, its count as its own report gives it; nothing "
       "learnt before it happened; the marks and scores run to the prototype's today",
       set(SOURCES) == stored and len(S) == len(p["sources"]) == len(SOURCES) + len(ESTIMATES)
       and set(reported) == set(SOURCES) and all(S[s]["records"] == n for s, n in reported.items()) and early == 0
       and all(x["learnt"] >= x["happened"] for x in S.values())
       and S["rdw_register"]["learnt"] == S["x4_synthetic_resale"]["learnt"] == read_on
       and S["marks"]["happened"] == S["readiness"]["happened"] == p["as_of"][:10],
       f"{len(SOURCES)} sources, each count from its report, and {len(ESTIMATES)} estimate tables; the register read "
       f"on {read_on}; learnt from {min(x['learnt'] for x in S.values())} to {max(x['learnt'] for x in S.values())}")
    ok("no VIN leaves the store", not re.search(r"\b[A-HJ-NPR-Z0-9]{17}\b", blob))
    ok("nothing at advert level (no listing ids or URLs)", "http" not in blob and "listing" not in blob.lower())
    ok("every group figure is traced", len(TRACE) > 0 and all(w for _, w in TRACE))
    return res


def main():
    con = duckdb.connect(str(STORE), read_only=True)
    BODY.update(bodies(con))
    g = group()
    p, claims, rd, arms = proto(con, g)
    # the tool pass, part 6: every listed car can be found and opened; each source's as-of date
    p["car_index"] = car_index(p)
    p["sources"] = source_dates(con)
    p["cars"] = cars(con, [r["car"] for r in p["car_index"]])
    data = {"built": str(date.today()), "group": g, "proto": p,
            "trace": sorted(set(TRACE), key=lambda t: (t[1], t[0]))}
    blob = json.dumps(data, ensure_ascii=False, default=str)
    res = checks(con, g, p, claims, rd, arms, blob)
    data["checks"] = [[n, c, d] for n, c, d in res]
    for n, c, d in res:
        print(("pass" if c else "FAIL") + f": {n}" + (f" ({d})" if d else ""))
    if not all(c for _, c, _ in res):
        raise SystemExit("export stopped: a check failed")
    OUT.write_text("window.DASH = " + json.dumps(data, ensure_ascii=False, default=str) + ";\n")
    print(f"wrote {OUT.name}: {OUT.stat().st_size / 1024:.0f} KB, {len(p['cars'])} cars in the drawer, "
          f"{len(data['trace'])} traced group figures")


if __name__ == "__main__":
    main()
