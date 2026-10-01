"""X8: the merger's diversification dividend. Does the merged book carry less level risk than PSA's and FCA's apart?

Part 1: each legacy book's country mix before the merger. The weights are new passenger-car registrations by Member
State and make in 2019 (the last year before COVID; 2018 as a check), from the EU's CO2 monitoring register
(Regulation (EU) 2019/631), which records every new car registered in the EU, Iceland and Norway, and the UK to 2020:
the EEA's DISCODATA SQL endpoint, final data (status F), summed over the registration count R. The legacy books:
  PSA  Peugeot, Citroen, DS, Opel, Vauxhall (PSA owned Opel and Vauxhall from 1 August 2017)
  FCA  Fiat, Alfa Romeo, Lancia, Jeep, Abarth, Maserati, and FCA's US makes
Registrations stand in for where each book's residuals sit. A finance arm's residual exposure by country would be
the better weight, and none is published (the user's research report, 26 September; X8's section).

Checks for part 1: every make of each group is found (none unmatched among the large makes); the Netherlands'
counts agree with the RDW register we hold (X2); shares sum to 1.

Usage: .venv/bin/python analysis/merger_diversification.py [--pull]   (writes analysis/merger_diversification_report.md;
       --pull queries the EEA again; the aggregated table is cached in data/reference/, private)
"""
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import requests

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE))
from build_unified import md_table  # noqa: E402

OUT = HERE / "merger_diversification_report.md"
CACHE = HERE.parent / "data" / "reference" / "x8_eea_registrations.parquet"
EEA = "https://discodata.eea.europa.eu/sql"
YEARS = (2018, 2019)
QUERY = ("SELECT Year, MS, Mh, Mk, SUM(R) AS n FROM [CO2Emission].[latest].[co2cars] "
         "WHERE Year = {year} AND Status = 'F' GROUP BY Year, MS, Mh, Mk")


def pull():
    """Registrations by year, Member State, manufacturer and make, aggregated on the EEA's server."""
    frames = []
    for year in YEARS:
        page = 1
        while True:
            r = requests.get(EEA, params={"query": QUERY.format(year=year), "p": page, "nrOfHits": 10000},
                             timeout=600)
            r.raise_for_status()
            rows = r.json()["results"]
            frames.append(pd.DataFrame(rows))
            if len(rows) < 10000:
                break
            page += 1
    d = pd.concat(frames, ignore_index=True)
    CACHE.parent.mkdir(parents=True, exist_ok=True)
    d.to_parquet(CACHE, index=False)
    return d


def registrations(refresh=False):
    return pull() if refresh or not CACHE.exists() else pd.read_parquet(CACHE)


GROUPS = {"PSA": {"PEUGEOT", "CITROEN", "DS", "OPEL", "VAUXHALL"},
          "FCA": {"FIAT", "ABARTH", "LANCIA", "ALFA", "JEEP", "MASERATI", "DODGE", "CHRYSLER", "RAM"}}
STEMS = ("PEUG", "CITRO", "OPEL", "VAUX", "FIAT", "ABART", "LANCI", "ALFA", "JEEP", "MASER", "DODGE", "CHRYS")
EEA_TO_PANEL = {"GR": "EL", "GB": "UK"}      # the register's codes against Eurostat's and ours
KBA = HERE.parent / "data" / "x2_kba_kurzzulassungen.csv"      # X2: KBA new registrations by brand and year
RDW = HERE.parent / "data" / "x2_rdw_brand_day.parquet"        # X2: RDW first registrations in NL by brand and day
REGISTER = HERE.parent / "assumptions.csv"
KBA_BRANDS = {"PSA": ("Peugeot", "Citroen", "DS", "Opel"), "FCA": ("Fiat", "Alfa Romeo", "Jeep")}
YEAR = 2019
SEED = 7


def norm(make):
    """A make as upper-case words, accents and punctuation removed."""
    import unicodedata
    s = unicodedata.normalize("NFKD", str(make)).encode("ascii", "ignore").decode().upper()
    return " ".join("".join(c if c.isalnum() else " " for c in s).split())


def family(make):
    words = set(norm(make).split())
    hits = [g for g, brands in GROUPS.items() if words & brands]
    return hits[0] if len(hits) == 1 else ("both" if hits else None)


def clean(d):
    """Drop the register's duplicate records and rows with no Member State; tag each make's legacy group."""
    d = d[(d["Mh"] != "DUPLICATE") & d["MS"].fillna("").ne("")].copy()
    d["market"] = d["MS"].replace(EEA_TO_PANEL)
    d["group"] = d["Mk"].map(family)
    return d


def mix(d, year=YEAR):
    """Registrations by market for each legacy group and the merged book, and each one's country shares."""
    y = d[(d["Year"] == year) & d["group"].isin(GROUPS)]
    n = y.pivot_table(index="market", columns="group", values="n", aggfunc="sum", fill_value=0)
    n["merged"] = n["PSA"] + n["FCA"]
    return n, n / n.sum()


def part1_checks(d, n):
    rows = []

    def check(name, passes, got):
        rows.append({"check": name, "got": got, "passes": bool(passes)})

    y = d[d["Year"] == YEAR]
    both = y.loc[y["group"] == "both", "n"].sum()
    loose = y[y["group"].isna() & y["Mk"].map(lambda m: any(s in norm(m) for s in STEMS))]
    check("no make is claimed by both groups", both == 0, f"{both} cars")
    check("no unassigned make carries a group brand's name (e.g. an accented or abbreviated spelling)",
          loose["n"].sum() == 0, ", ".join(sorted(set(loose["Mk"]))) or "none")
    check("each group's country shares sum to 1", np.allclose((n / n.sum()).sum(), 1.0), "PSA, FCA, merged")

    # KBA counts motorhomes as passenger cars under the base vehicle's make; the CO2 register excludes them as
    # special-purpose vehicles. So Germany is compared with KBA less its motorhomes (register rows kba_wohnmobile_*).
    reg = pd.read_csv(REGISTER).set_index("id")["value"]
    kba = pd.read_csv(KBA)
    kba = kba[kba["year"] == YEAR].set_index("brand")["new"]
    de = y[y["market"] == "DE"]
    homes = {"total": reg["kba_wohnmobile_2019"], "PSA": reg["kba_wohnmobile_citroen_jumper_2019"],
             "FCA": reg["kba_wohnmobile_fiat_ducato_2019"]}
    raw_gap = de["n"].sum() / kba["Insgesamt"] - 1
    gap = de["n"].sum() / (kba["Insgesamt"] - homes["total"]) - 1
    check(f"Germany: the register's 2019 total within 1% of KBA's less its motorhomes (against all of KBA's: "
          f"{raw_gap:+.2%})", abs(gap) <= 0.01, f"{gap:+.2%}")
    for g, brands in KBA_BRANDS.items():
        eea = de[de["Mk"].map(lambda m: norm(m).split()[0] if norm(m) else "").isin(
            [norm(b).split()[0] for b in brands])]["n"].sum()
        raw_gap = eea / kba[list(brands)].sum() - 1
        gap = eea / (kba[list(brands)].sum() - homes[g]) - 1
        base = "Fiat Ducato" if g == "FCA" else "Citroen Jumper (Peugeot and Opel bases not listed)"
        check(f"Germany, {g} brands KBA lists ({', '.join(brands)}): the register within 3% of KBA less motorhomes "
              f"on the {base} (against all of KBA's: {raw_gap:+.2%})", abs(gap) <= 0.03, f"{gap:+.2%}")

    rdw = pd.read_parquet(RDW)
    rdw = rdw[pd.to_datetime(rdw["d"]).dt.year == YEAR]
    rdw_n = rdw.assign(group=rdw["merk"].map(family)).groupby("group")["n"].sum()
    nl = y[y["market"] == "NL"]
    gap = nl["n"].sum() / rdw["n"].sum() - 1
    check("Netherlands: the register's 2019 total within 3% of RDW's first registrations", abs(gap) <= 0.03,
          f"{gap:+.2%}")
    for g in GROUPS:
        gap = nl.loc[nl["group"] == g, "n"].sum() / rdw_n.get(g, np.nan) - 1
        check(f"Netherlands, {g}: the register within 5% of RDW", abs(gap) <= 0.05, f"{gap:+.2%}")
    return pd.DataFrame(rows)


def level_markets():
    """The markets with a level series over X6's common window: its 26 Eurostat markets, and the UK from the ONS."""
    from level_charge import common_panel
    common, core, _ = common_panel()
    return sorted(common) + ["UK"], core


def concentration(shares):
    """Herfindahl index of a book's country shares, and its inverse, the effective number of markets."""
    h = float((shares ** 2).sum())
    return h, 1 / h


def part1(d):
    n, s = mix(d)
    n18, s18 = mix(d, 2018)
    covered, core = level_markets()
    inside = s.index.isin(covered)
    cover = s[inside].sum()
    table = n.join(s, rsuffix=" share").sort_values("merged", ascending=False)
    table["level series"] = ["ONS (UK)" if m == "UK" else ("Eurostat" if m in covered else "none")
                             for m in table.index]
    show = pd.DataFrame({"market": table.index,
                         "PSA cars": table["PSA"].map("{:,.0f}".format),
                         "PSA share": table["PSA share"].map("{:.1%}".format),
                         "FCA cars": table["FCA"].map("{:,.0f}".format),
                         "FCA share": table["FCA share"].map("{:.1%}".format),
                         "merged share": table["merged share"].map("{:.1%}".format),
                         "level series": table["level series"]})
    conc = pd.DataFrame([{"book": g, "cars, 2019": f"{n[g].sum():,.0f}",
                          "share in markets with a level series": f"{cover[g]:.1%}",
                          "Herfindahl index": f"{concentration(s[g])[0]:.3f}",
                          "effective number of markets": f"{concentration(s[g])[1]:.1f}",
                          "largest market": f"{s[g].idxmax()} {s[g].max():.1%}",
                          "largest share change, 2018 to 2019 (points)":
                              f"{(s[g] - s18[g].reindex(s.index).fillna(0)).abs().max() * 100:.1f}"}
                         for g in ("PSA", "FCA", "merged")])
    return n, s, show, conc, cover, covered, core


def part1_text(show, conc, ck1):
    return [
        "## Part 1: each legacy book's country mix before the merger",
        "",
        f"New passenger cars registered in {YEAR}, by market, from the EU's CO2 monitoring register (EEA DISCODATA, "
        "final data, duplicates removed). PSA includes Opel and Vauxhall. The markets with a level series are X6's "
        "26 Eurostat markets over its common window and the UK (ONS); Belgium's Eurostat series starts in December "
        "2018, after that window, so it has none here (part 2 adds it on a shorter window).",
        "",
        md_table(conc),
        "",
        "The effective number of markets is one over the Herfindahl index: the number of equal markets that would be "
        "as concentrated. It counts spread, not risk; part 2 measures risk.",
        "",
        md_table(show),
        "",
        "### Part 1 checks",
        "",
        md_table(ck1),
        "",
    ]



# ---------- Part 2: the dividend ----------

ONS_SERIES = "ONS CPI 07.1.1B, second-hand cars (D7E9)"
CURRENCY = {"UK": "GBP", "PL": "PLN", "CZ": "CZK", "HU": "HUF", "SE": "SEK", "DK": "DKK", "RO": "RON", "BG": "BGN"}
HORIZONS = (12, 36, 48)
BELGIUM_START = "2018-12"      # Belgium's Eurostat series starts here


def series_window(start, end, france=None, euros=False):
    """Each market's monthly level over [start, end]: Eurostat, the UK from the ONS; only complete series kept."""
    from level_risk import EUROSTAT, load
    raw = load()
    s = {g: v for (name, g), v in raw.items() if name == EUROSTAT}
    s["UK"] = raw[(ONS_SERIES, "UK")]
    if france:
        s["FR"] = s[france]
    months = pd.period_range(start, end, freq="M").astype(str)
    out = {}
    for g, v in s.items():
        v = v.reindex(months)
        if v.notna().all():
            out[g] = v
    if euros:
        fx = pd.read_csv(HERE.parent / "data" / "raw" / "fx" / "ecb_monthly.csv",
                         usecols=["CURRENCY", "TIME_PERIOD", "OBS_VALUE"])
        for g, cur in CURRENCY.items():
            if g in out:
                rate = fx[fx["CURRENCY"] == cur].set_index("TIME_PERIOD")["OBS_VALUE"].reindex(months)
                assert rate.notna().all(), (g, cur)
                out[g] = out[g] / rate          # the level in euros: local index over units per euro
    return out


def market_moves(series, k):
    """Every market's move over k months from each start month, as fractions; rows are start months."""
    return pd.DataFrame({g: (v.shift(-k) / v - 1) for g, v in series.items()}).dropna()


def book_weights(n, markets):
    """Each book's shares over the markets with a series, and each legacy book's share of the merged cars there."""
    m = n.loc[n.index.isin(markets)]
    w = m / m.sum()
    size = m[["PSA", "FCA"]].sum() / m["merged"].sum()
    return w, size


def tail(x, share=0.1):
    """Mean of the worst tenth of a loss series (largest losses)."""
    x = np.sort(np.asarray(x, dtype=float))[::-1]
    return float(x[: int(np.ceil(share * len(x)))].mean())


def measures(series, n, strike_markets=None):
    """For PSA, FCA and merged: the one-year p10 and worst move of the book, and the 36- and 48-month loss per euro of
    residual (mean and worst tenth) at X6's two strikes. The loss of a book in a window is the weighted sum of its
    markets' shortfalls, each car keeping its own market's residual."""
    from level_charge import shortfall
    markets = [g for g in series if g in n.index and n.loc[g, "merged"] > 0]
    w, size = book_weights(n, markets)
    out = {"weights": w, "size": size, "markets": markets}
    one = market_moves({g: series[g] for g in markets}, 12)
    for b in ("PSA", "FCA", "merged"):
        mv = one[markets] @ w.loc[markets, b]
        out[(b, 12, "p10")] = float(np.percentile(mv, 10))
        out[(b, 12, "worst")] = float(mv.min())
        out[(b, 12, "series")] = mv
    for k in (36, 48):
        mk = market_moves({g: series[g] for g in markets}, k)
        for b in ("PSA", "FCA", "merged"):
            out[(b, k, "p10")] = float(np.percentile(mk[markets] @ w.loc[markets, b], 10))
        pooled = mk[strike_markets or markets].to_numpy().ravel()
        for strike, assumed in (("median", float(np.median(pooled))), ("no change", 0.0)):
            loss = pd.DataFrame({g: shortfall(mk[g], assumed) for g in markets}, index=mk.index)
            for b in ("PSA", "FCA", "merged"):
                lb = loss[markets] @ w.loc[markets, b]
                out[(b, k, strike, "mean")] = float(lb.mean())
                out[(b, k, strike, "tail")] = tail(lb)
            out[(k, strike, "assumed")] = assumed
    return out


def dividends(m):
    """Apart (each legacy book's figure, weighted by its share of the merged cars) against merged, per measure."""
    sP, sF = m["size"]["PSA"], m["size"]["FCA"]
    rows = []
    for key, label in (((12, "p10"), "one-year fall, p10"),
                       ((12, "worst"), "one-year fall, worst")):
        loss = {b: max(0.0, -m[(b, *key)]) for b in ("PSA", "FCA", "merged")}
        apart = sP * loss["PSA"] + sF * loss["FCA"]
        rows.append({"measure": label, "PSA": loss["PSA"], "FCA": loss["FCA"], "apart": apart,
                     "merged": loss["merged"], "dividend": apart - loss["merged"]})
    for k in (36, 48):
        for strike in ("median", "no change"):
            for stat, name in (("tail", "worst tenth"), ("mean", "expected")):
                v = {b: m[(b, k, strike, stat)] for b in ("PSA", "FCA", "merged")}
                apart = sP * v["PSA"] + sF * v["FCA"]
                rows.append({"measure": f"{k}-month loss per euro of residual, {name}, {strike} strike",
                             "PSA": v["PSA"], "FCA": v["FCA"], "apart": apart, "merged": v["merged"],
                             "dividend": apart - v["merged"]})
    d = pd.DataFrame(rows)
    d["dividend, % of apart"] = np.where(d["apart"] > 0, d["dividend"] / d["apart"], np.nan)
    return d


def part2_checks(series, n, m):
    """Reproduction, bookkeeping, and the identities that must hold exactly."""
    from level_risk import pooled
    from level_charge import common_panel
    rows = []

    def check(name, passes, got):
        rows.append({"check": name, "got": got, "passes": bool(passes)})

    reg = pd.read_csv(REGISTER).set_index("id")["value"]
    common, core, _ = common_panel()
    for rid, geos in (("level_1y_p10_core", core), ("level_1y_p10", sorted(common))):
        got = float(np.percentile(pooled(common, geos, 12), 10))
        check(f"the workbook's {rid} reproduced from the same pooled one-year windows ({reg[rid]:+.1f}%)",
              round(got, 1) == reg[rid], f"{got:+.2f}%")
    w, size = m["weights"], m["size"]
    ident = (size["PSA"] * w["PSA"] + size["FCA"] * w["FCA"] - w["merged"]).abs().max()
    check("the merged weights are exactly the legacy books' weights mixed by their sizes", ident < 1e-15,
          f"largest gap {ident:.1e}")
    gaps = [abs(m[("merged", k, st, "mean")] - size["PSA"] * m[("PSA", k, st, "mean")]
                - size["FCA"] * m[("FCA", k, st, "mean")]) for k in (36, 48) for st in ("median", "no change")]
    check("the expected loss adds up: merged equals the books apart, every horizon and strike", max(gaps) < 1e-15,
          f"largest gap {max(gaps):.1e}")
    over = [m[("merged", k, st, "tail")] - size["PSA"] * m[("PSA", k, st, "tail")]
            - size["FCA"] * m[("FCA", k, st, "tail")] for k in (36, 48) for st in ("median", "no change")]
    check("the worst tenth is subadditive: merged never above the books apart", max(over) <= 1e-15,
          f"largest excess {max(over):.1e}")
    mk = market_moves({g: series[g] for g in m["markets"]}, 12)
    origin = mk.index[len(mk) // 2]
    by_hand = sum(w.loc[g, "merged"] * (series[g].iloc[list(series[g].index).index(origin) + 12]
                                         / series[g][origin] - 1) for g in m["markets"])
    got = float(m[("merged", 12, "series")][origin])
    check(f"bookkeeping: the merged book's one-year move from {origin} recomputed market by market",
          abs(by_hand - got) < 1e-12, f"{got:+.4%}")

    months = pd.period_range("2000-01", periods=120, freq="M").astype(str)
    rng = np.random.default_rng(SEED)
    walk = pd.Series(np.exp(np.cumsum(rng.normal(0, 0.01, 120))), index=months)
    same = pd.DataFrame({"PSA": [1.0], "FCA": [1.0], "merged": [2.0]}, index=["A"])
    toy = measures({"A": walk}, same)
    div = dividends(toy)["dividend"].abs().max()
    check("toy: both books in the same single market give no dividend on any measure", div < 1e-15, f"{div:.1e}")
    mirror = pd.Series(1 / walk.to_numpy(), index=months)
    split = pd.DataFrame({"PSA": [1.0, 0.0], "FCA": [0.0, 1.0], "merged": [1.0, 1.0]}, index=["A", "B"])
    toy = measures({"A": walk, "B": mirror}, split)
    one = [abs(toy[("merged", 12, "series")]).max()]
    check("toy: two books in mirror-image markets (log moves opposite) merge to a book that barely moves",
          one[0] < 0.01, f"largest one-year move {one[0]:.2%}")
    return pd.DataFrame(rows)


VARIANTS = {
    "main: 2019 weights, local currency, Dec 2016-Dec 2025": {},
    "without France": {"drop": ["FR"]},
    "France replaced by the euro-area index (EA20, which contains France)": {"france": "EA20"},
    "with Belgium, Dec 2018-Dec 2025": {"start": BELGIUM_START},
    "the same shorter window without Belgium": {"start": BELGIUM_START, "drop": ["BE"]},
    "in euros (non-euro markets converted at ECB rates)": {"euros": True},
    "2018 weights": {"year": 2018},
}


def run_variant(d, spec):
    from level_risk import COMMON_END, COMMON_START
    series = series_window(spec.get("start", COMMON_START), COMMON_END, spec.get("france"), spec.get("euros", False))
    for g in spec.get("drop", []):
        series.pop(g, None)
    n, _ = mix(d, spec.get("year", YEAR))
    return series, n, measures(series, n)


def comovement(series, m):
    """How closely the two legacy books move. Correlation over all one-year windows, and how often their worst tenth of
    windows coincide at each horizon (conditioning a correlation on the merged book's bad windows would bias it
    negative, since selecting on a low sum forces the parts apart)."""
    from level_charge import shortfall
    rows = []
    a, b = m[("PSA", 12, "series")], m[("FCA", 12, "series")]
    rows.append({"measure": "correlation of the two books' one-year moves, all windows", "value": f"{a.corr(b):.2f}",
                 "windows": len(a)})
    mk = market_moves({g: series[g] for g in ("FR", "IT", "DE", "ES", "UK")}, 12)
    pairs = mk.corr().where(np.triu(np.ones((5, 5), bool), 1)).stack()
    rows.append({"measure": "median correlation of one-year moves between FR, IT, DE, ES and UK",
                 "value": f"{pairs.median():.2f}", "windows": len(mk)})
    rows.append({"measure": "correlation of France's and Italy's one-year moves", "value": f"{mk['FR'].corr(mk['IT']):.2f}",
                 "windows": len(mk)})
    w = m["weights"]
    for k in HORIZONS:
        mv = market_moves({g: series[g] for g in m["markets"]}, k)
        if k == 12:
            loss = {bk: -(mv[m["markets"]] @ w.loc[m["markets"], bk]) for bk in ("PSA", "FCA")}
        else:
            assumed = m[(k, "median", "assumed")]
            sf = pd.DataFrame({g: shortfall(mv[g], assumed) for g in m["markets"]}, index=mv.index)
            loss = {bk: sf[m["markets"]] @ w.loc[m["markets"], bk] for bk in ("PSA", "FCA")}
        size = int(np.ceil(0.1 * len(mv)))
        worst = {bk: set(v.sort_values(ascending=False).index[:size]) for bk, v in loss.items()}
        rows.append({"measure": f"share of the two books' worst-tenth {k}-month windows that coincide",
                     "value": f"{len(worst['PSA'] & worst['FCA']) / size:.0%}", "windows": len(mv)})
    return pd.DataFrame(rows)


def core_tail(k=12):
    """Which core markets supply the worst tenth of the pooled one-year windows behind the workbook's p10."""
    from level_charge import common_panel
    common, core, _ = common_panel()
    mv = pd.concat([pd.DataFrame({"market": g, "move": (common[g].shift(-k) / common[g] - 1).dropna()}) for g in core])
    cut = mv["move"].quantile(0.1)
    share = mv.loc[mv["move"] <= cut, "market"].value_counts(normalize=True).reindex(core).fillna(0)
    return pd.DataFrame({"market": share.index, "share of the pooled worst tenth": share.map("{:.0%}".format),
                         "the market's own p10 one-year move": [f"{np.percentile(mv.loc[mv.market == g, 'move'], 10):+.1%}"
                                                                for g in share.index]})


def pct(v, digits=2):
    """A fraction as a percentage; float noise around zero prints as 0."""
    v = 100 * v
    return f"{0.0 if abs(v) < 0.5 * 10 ** -digits else v:.{digits}f}"


def variant_table(d):
    rows = []
    for name, spec in VARIANTS.items():
        series, n, m = run_variant(d, spec)
        dv = dividends(m).set_index("measure")
        one, worst = dv.loc["one-year fall, p10"], dv.loc["one-year fall, worst"]
        rows.append({"variant": name, "markets": len(m["markets"]),
                     "one-year p10 fall, apart (%)": pct(one["apart"]), "one-year p10 fall, merged (%)": pct(one["merged"]),
                     "one-year p10 dividend (% of apart)": pct(one["dividend, % of apart"], 0),
                     "worst one-year fall, merged (%)": pct(worst["merged"]),
                     "worst one-year dividend (% of apart)": pct(worst["dividend, % of apart"], 0),
                     "36-month worst-tenth dividend (% of apart)":
                         pct(dv.loc["36-month loss per euro of residual, worst tenth, median strike", "dividend, % of apart"], 1),
                     "48-month worst-tenth dividend (% of apart)":
                         pct(dv.loc["48-month loss per euro of residual, worst tenth, median strike", "dividend, % of apart"], 1)})
    return pd.DataFrame(rows)


def part2_text(m, dv, variants, como, ck2, tail_src):
    reg = pd.read_csv(REGISTER).set_index("id")["value"]
    show = dv.copy()
    for c in ("PSA", "FCA", "apart", "merged", "dividend"):
        show[c] = show[c].map(lambda v: pct(v))
    show["dividend, % of apart"] = dv["dividend, % of apart"].map(lambda v: "" if pd.isna(v) else pct(v, 1))
    show = show.rename(columns={"PSA": "PSA (%)", "FCA": "FCA (%)", "apart": "apart (%)", "merged": "merged (%)",
                                "dividend": "dividend (points)"})
    one = dv.set_index("measure").loc["one-year fall, p10"]
    worst = dv.set_index("measure").loc["one-year fall, worst"]
    spread = pd.DataFrame([
        {"figure": "one core market's p10 one-year fall, the 7 core markets pooled (level_1y_p10_core; leak 3's stress)",
         "fall (%)": f"{abs(reg['level_1y_p10_core']):.1f}"},
        {"figure": "the merged book's p10 one-year fall, spread as its 2019 registrations",
         "fall (%)": pct(one["merged"])},
        {"figure": "the worst one-year fall in any of the 26 markets (level_1y_worst; a reference line in no total)",
         "fall (%)": f"{abs(reg['level_1y_worst']):.1f}"},
        {"figure": "the merged book's worst one-year fall", "fall (%)": pct(worst["merged"])}])
    size = m["size"]
    return [
        "## Part 2: the dividend",
        "",
        "A book's move over a window is the sum of its markets' moves weighted by its registrations (part 1), over the "
        "markets with a series. A book's loss per euro of residual is the same weighted sum of its markets' shortfalls, "
        "each car keeping its own market's residual (X6's `shortfall`, at X6's median strike, the pooled median move, "
        "or no change). **Apart** is PSA's and FCA's own figures weighted by their shares of the merged cars "
        f"({size['PSA']:.1%} and {size['FCA']:.1%}); **merged** is the merged book's. The dividend is apart less merged. "
        "The expected loss cannot move: each car's loss depends only on its own market, so expected losses add (checked). "
        "Only the tail can: the p10 and worst one-year fall that leak 3 uses, and X6's worst tenth at 36 and 48 months.",
        "",
        "### The main run: 2019 weights, local currency, December 2016 to December 2025",
        "",
        md_table(show),
        "",
        "### Sensitivities",
        "",
        "France's official index barely moves (STATE data trap), so France is dropped, or replaced by the euro-area "
        "index, which still contains it. Belgium enters on the shorter window its series allows, beside the same window "
        "without it. In euros, non-euro markets' levels are converted at ECB monthly rates. 2018 weights check the year.",
        "",
        md_table(variants),
        "",
        "### Why: how the two books move",
        "",
        md_table(como),
        "",
        "### The spread book against a single market",
        "",
        "One market's one-in-ten-year fall, applied to the whole buy-back book, treats it as if it sat in one country. "
        "A book spread like the group's has a milder bad year on the same index, and the pooled core-market p10 comes "
        "mostly from one market:",
        "",
        md_table(spread),
        "",
        md_table(tail_src),
        "",
        f"On the index, the spread book's bad year is well below the single-market figure. X7 found the prices a lessor "
        f"meets move about {reg['x7_ald_slope_36m']:.1f} times the index (`x7_ald_slope_36m`); even scaled by that, the "
        f"merged book's p10 fall would be {100 * one['merged'] * reg['x7_ald_slope_36m']:.1f}%, against the "
        f"single market's {abs(reg['level_1y_p10_core']):.1f}%. So leak 3's base is the book's own fall, and the "
        "single-market figure is its stress (part 3).",
        "",
        "### Part 2 checks",
        "",
        md_table(ck2),
        "",
    ]


# ---------- Part 3: in euros ----------

WORKBOOK = HERE.parent.parent / "Case4_Value_at_Risk.xlsx"
LEAK3 = "Leak 3 - Residual value"


BASE_VARIANT = "France replaced by the euro-area index (EA20, which contains France)"


def workbook_value(prefix):
    """A leak 3 line's stored value, found by its label; None when no single line carries it."""
    from openpyxl import load_workbook
    ws = load_workbook(WORKBOOK, data_only=True)[LEAK3]
    hits = [row[1].value for row in ws.iter_rows() if isinstance(row[0].value, str)
            and row[0].value.strip().startswith(prefix)]
    return float(hits[0]) if len(hits) == 1 else None


def part3(d, m, dv):
    reg = pd.read_csv(REGISTER).set_index("id")["value"]
    book = reg["buyback_payables_current_eur_m"]
    pooling = (book * reg["curve_1y_p80_unknown"] / 2 / reg["retained_1y_pooled"] * reg["thin_share_of_book"] / 100
               * reg["pooling_gain_100"] / 100)
    _, _, mb = run_variant(d, VARIANTS[BASE_VARIANT])
    move_1y, move_36 = mb[("merged", 12, "p10")], mb[("merged", 36, "p10")]
    core, worst_any = book * abs(reg["level_1y_p10_core"]) / 100, book * abs(reg["level_1y_worst"]) / 100
    rows = []

    def check(name, passes, got):
        rows.append({"check": name, "got": got, "passes": bool(passes)})

    base_reg = book * abs(reg.get("level_1y_p10_book", np.nan)) / 100
    for label, mine, prefix in (("B8's pooling line", pooling, "Pooling the group's markets"),
                                ("leak 3's base level shock (the book's own fall)", base_reg, "Base: the group's book, spread"),
                                ("leak 3's stress level shock (a core market's fall)", core, "Stress: the whole book in a core market's bad year"),
                                ("the reference line (the worst fall in any market)", worst_any,
                                 "Reference, in no total: the worst twelve-month")):
        stored = workbook_value(prefix)
        check(f"{label} recomputed from the register equals the workbook's stored value "
              f"({'missing' if stored is None else f'{stored:,.3f}'})",
              stored is not None and abs(mine - stored) < 1e-6, f"{mine:,.3f}")
    check("the register's level_1y_p10_book is this run's figure, to the register's three decimals",
          abs(reg.get("level_1y_p10_book", np.nan) - round(100 * move_1y, 3)) < 1e-9,
          f"{100 * move_1y:+.4f}%")
    ck3 = pd.DataFrame(rows)

    div = dv.set_index("measure")
    one, worst = div.loc["one-year fall, p10"], div.loc["one-year fall, worst"]
    spans = []
    for name, spec in VARIANTS.items():
        _, _, mv = run_variant(d, spec)
        dd = dividends(mv).set_index("measure")
        spans.append({"variant": name,
                      "base dividend": book * dd.loc["one-year fall, p10", "dividend"],
                      "merged base": book * dd.loc["one-year fall, p10", "merged"]})
    spans = pd.DataFrame(spans)
    nofr = spans.set_index("variant").loc["without France", "merged base"]
    slope = reg["x7_ald_slope_36m"]
    inputs = pd.DataFrame([
        {"figure": "the merged book's one-year move, p10, France on the euro-area index (leak 3's base)",
         "move (%)": f"{100 * move_1y:+.3f}"},
        {"figure": "the merged book's 36-month move, p10, France on the euro-area index (the three-year context's base)",
         "move (%)": f"{100 * move_36:+.3f}"}])
    from level_risk import load
    uk = load()[(ONS_SERIES, "UK")].dropna()
    uk1 = (uk.shift(-12) / uk - 1).dropna()
    stress_check = pd.DataFrame([
        {"figure": "a core market's one-year move, p10, pooled (level_1y_p10_core; leak 3's stress)",
         "move (%)": f"{reg['level_1y_p10_core']:+.1f}"},
        {"figure": f"the UK market's one-year move, p10, ONS {uk.index.min()[:4]}-{uk.index.max()[:4]} (several cycles)",
         "move (%)": f"{100 * np.percentile(uk1, 10):+.2f}"},
        {"figure": f"the UK market's worst one-year move, ONS, from {uk1.idxmin()}",
         "move (%)": f"{100 * uk1.min():+.2f}"}])
    euros = pd.DataFrame([
        {"figure": "leak 3's base level shock: the merged book's own p10 fall, France on the euro-area index "
                   "(the workbook)", "EUR m a year": f"{book * abs(move_1y):.1f}"},
        {"figure": "the same, France's own index", "EUR m a year": f"{book * one['merged']:.1f}"},
        {"figure": "the same without France", "EUR m a year": f"{nofr:.1f}"},
        {"figure": f"leak 3's base scaled by X7's lessor-to-index slope ({slope:.2f}, x7_ald_slope_36m)",
         "EUR m a year": f"{book * abs(move_1y) * slope:.1f}"},
        {"figure": "leak 3's stress level shock: the whole book at a core market's p10 fall (the workbook)",
         "EUR m a year": f"{core:.1f}"},
        {"figure": "the worst one-year fall in any market, on the whole book (a reference line in no total)",
         "EUR m a year": f"{worst_any:.1f}"},
        {"figure": "the merged book's own worst fall, France's own index (not the stress: its history holds no "
                   "Europe-wide crash)", "EUR m a year": f"{book * worst['merged']:.1f}"},
        {"figure": "the merger's dividend on leak 3's base: each legacy book's own one-in-ten-year fall against the "
                   "merged book's (main run)", "EUR m a year": f"{book * one['dividend']:.1f}"},
        {"figure": "the same, lowest across the sensitivities", "EUR m a year": f"{spans['base dividend'].min():.1f}"},
        {"figure": "the same, highest across the sensitivities", "EUR m a year": f"{spans['base dividend'].max():.1f}"},
        {"figure": "the merger's dividend on each book's own worst year, on leak 3's book (main run)",
         "EUR m a year": f"{book * worst['dividend']:.1f}"},
        {"figure": "B8: pooling the group's markets, on the thin slice (the workbook)", "EUR m a year": f"{pooling:.1f}"}])
    price = reg["eu_revenue_eur_m"] * 1e6 / reg["eu_shipments"]
    kept = {36: reg["retained_3y_uk"] / 100, 48: reg["retained_4y_uk"] / 100}
    contract = pd.DataFrame([{"figure": f"the merger's dividend per {k}-month contract in the worst tenth of windows, "
                                        f"X6's group car ({strike} strike)",
                              "EUR per contract": f"{price * kept[k] * div.loc[f'{k}-month loss per euro of residual, worst tenth, {strike} strike', 'dividend']:.0f}"}
                             for k in (36, 48) for strike in ("median", "no change")])
    return inputs, stress_check, euros, contract, ck3


def part3_text(inputs, stress_check, euros, contract, ck3):
    return [
        "## Part 3: in euros, and leak 3's level shock",
        "",
        "Leak 3's book is the group's buy-back payables due within a year (`buyback_payables_current_eur_m`), the "
        "group's own balance sheet: unlike the pre-merger finance joint ventures, whose books sat with different "
        "partners, this is where a diversification gain would accrue to the group. The legacy books are split by their "
        "2019 registrations (part 1), a proxy: the buy-back book's own country mix is not published.",
        "",
        "**Leak 3's level shock (decided 26 September).** The base is the merged book's own one-in-ten-year fall. "
        "France is carried by the euro-area index, which contains it, because France's own index is smoothed guide "
        "values that barely move (STATE data trap); the main run and the run without France bracket it. Local "
        "currency, not euros: a market's fall in its own currency is what opens the gap between the car and the "
        "buy-back price, and a currency move shifts both. Both base figures are floors: the index understates the "
        "level a lessor meets (X7).",
        "",
        md_table(inputs),
        "",
        "**The stress** is a Europe-wide bad year as deep as a core market's one-in-ten: the whole book at "
        "`level_1y_p10_core`. The book's own worst year cannot be the stress: its history is about nine independent "
        "years with no Europe-wide crash (the Eurostat series for Germany, Italy and Spain start in 2014-15). And bad "
        "periods are shared: the two legacy books' worst-tenth 48-month windows coincide (part 2). The stress is not "
        "extreme. The UK market as a whole, a spread of regions and makes like a book, fell further than it one year "
        "in ten over several cycles:",
        "",
        md_table(stress_check),
        "",
        md_table(euros),
        "",
        "Per contract, at X6's group car (list price from EU revenue over shipments, times value retained):",
        "",
        md_table(contract),
        "",
        "**Reading it.** The merger's diversification is worth little at the one-year horizon and nothing at the "
        "contract horizons that price a lease. B8's pooling figure measures a different thing (better models from "
        "borrowing other markets' data), so the two are not substitutes; neither is large. The larger number here is "
        "not the merger's: the book was always spread across countries, which leak 3's base now credits.",
        "",
        "### Part 3 checks",
        "",
        md_table(ck3),
        "",
        "## Limits",
        "",
        "- **Registrations stand in for the books.** No finance arm publishes its residual exposure by country; the "
        "buy-back book (its buyers not published beyond the group's own leasing and rental joint ventures; most of it due "
        "within a year, a rental fleet's term) may sit differently from retail registrations.",
        "- **One window, one cycle.** December 2016 to December 2025 holds one rise and one fall; the one-year tail "
        "rests on about a dozen windows and the 48-month worst tenth on a handful, all overlapping.",
        "- **The index is not what a lessor meets.** Official indices understate the level a lessor meets (X7, part 3 "
        "and 3b), and France's barely moves; the sensitivities bound the France effect, not the basis gap.",
        "- **Registrations are counts, not euros.** Residual values per car differ by market; weights by value would "
        "tilt towards dearer markets.",
        "- **Capital:** CRR 134(7)'s standardised charge is per leased asset (EBA Q&A 2016_3072), so none of this "
        "lowers regulatory capital; it lowers an internal (economic) measure of the tail only.",
        "",
    ]

def main():
    refresh = "--pull" in sys.argv
    d = clean(registrations(refresh))
    n, s, show, conc, cover, covered, core = part1(d)
    ck1 = part1_checks(d, n)
    series, n_main, m = run_variant(d, {})
    dv = dividends(m)
    ck2 = part2_checks(series, n_main, m)
    variants = variant_table(d)
    como = comovement(series, m)
    tail_src = core_tail()
    lines = ["# X8: the merger's diversification dividend", "",
             "Generated by `analysis/merger_diversification.py`; the method is in its docstring.", ""]
    lines += part1_text(show, conc, ck1)
    lines += part2_text(m, dv, variants, como, ck2, tail_src)
    inputs, stress_check, euros, contract, ck3 = part3(d, m, dv)
    lines += part3_text(inputs, stress_check, euros, contract, ck3)
    OUT.write_text("\n".join(lines) + "\n")
    print(pd.concat([ck1, ck2, ck3])[["got", "passes"]].to_string(index=False))
    print(inputs.to_string(index=False))
    print(stress_check.to_string(index=False))
    print(euros.to_string(index=False))
    print(contract.to_string(index=False))
    print(variants.to_string(index=False))
    print(f"wrote {OUT.relative_to(HERE.parent)}")


if __name__ == "__main__":
    main()
